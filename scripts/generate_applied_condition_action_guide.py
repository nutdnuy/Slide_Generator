from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE_XLSX = Path("/Users/nuthdanai/Downloads/AI_Risk_Assessment_v2and_Minimum Control Checklist.xlsx")
SOURCE_PPTX = Path("/Users/nuthdanai/Downloads/AISLC.pptx")
PROJECT = "aislc-applied-condition-action-guide"
OUT_PPTX = ROOT / "outputs" / f"{PROJECT}.pptx"
OUTLINE = ROOT / "notes" / f"{PROJECT}-outline.md"
PLAN_JSON = ROOT / "notes" / f"{PROJECT}-plan.json"

SLIDE_W = 13.333
SLIDE_H = 7.5

FONT_HEAD = "Krungsri Simple"
FONT_BODY = "Sarabun"

C = {
    "white": "FFFFFF",
    "paper": "FBFAF8",
    "ink": "504545",
    "muted": "77706E",
    "taupe": "37302E",
    "brown": "6F625F",
    "line": "D9D3D0",
    "soft": "F3F0EE",
    "soft_yellow": "FEFAEE",
    "yellow": "FFC800",
    "yellow2": "FFD401",
    "green": "2E7D32",
    "red": "B3261E",
    "blue": "315E7D",
    "gray": "A6A6A6",
}


def rgb(hex_value: str) -> RGBColor:
    hex_value = hex_value.strip("#")
    return RGBColor(int(hex_value[0:2], 16), int(hex_value[2:4], 16), int(hex_value[4:6], 16))


def clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\u200b", " ")).strip()


def normalize_condition(value) -> str:
    text = clean(value)
    if not text:
        return ""
    if text == "เป็นโมเดลที่ใช้ \"ตัดสินใจในทางธุรกิจ การสนับสนุนการปฏิบัติงานหรือระบบงานสำคัญ และการให้บริการแก่ลูกค้า\" เช่น มีความเกี่ยวข้องกับ credit, risk, compliance เป็นต้น เป็น Supervised Learning":
        return "Business-critical supervised model"
    if text.startswith("เป็นโมเดลที่ใช้"):
        return "Business-critical supervised model"
    return text


def read_workbook():
    wb = load_workbook(SOURCE_XLSX, data_only=True)
    guideline = wb["Control_Guideline"]
    headers = [c.value for c in guideline[1]]
    index = {h: i for i, h in enumerate(headers) if h}
    rows = []
    for row_idx, raw in enumerate(guideline.iter_rows(min_row=2, values_only=True), start=2):
        if not any(raw):
            continue
        action = clean(raw[index["Action สิ่งที่โปรเจคทีมต้องทำ"]])
        condition = normalize_condition(raw[index["Applied condition for project"]])
        if not action:
            continue
        if not condition and clean(raw[index["Domain Area"]]) == "Model from 3rd Party":
            condition = "ใช้ Model จาก Vendor / 3rd Party"
        rows.append(
            {
                "row": row_idx,
                "domain": clean(raw[index["Domain Area"]]),
                "owner": clean(raw[index["Control Owner"]]),
                "condition": condition,
                "action": action,
                "evidence": clean(raw[index["Evident"]])
                or clean(raw[index["Evidence / Artifact"]])
                or clean(raw[index["Deliverable / Artifact ใน Pocedure"]]),
                "applicable": raw[index["Applicable"]] is True,
                "phase": clean(raw[index["AI SLC phase"]]),
                "gate": clean(raw[index["Gate Controller"]]),
                "process": clean(raw[index["Process Name"]]),
            }
        )

    status_sheet = wb["Applied condition for project"]
    condition_status = {}
    for raw in status_sheet.iter_rows(min_row=2, values_only=True):
        condition = normalize_condition(raw[0])
        if not condition or condition.startswith("จำนวน") or condition.startswith("วิธีใช้งาน") or condition[0:1].isdigit():
            continue
        status = raw[1]
        if status in ("Yes", "No"):
            condition_status[condition] = status
        elif status is not None:
            condition_status[condition] = str(status)

    grouped = defaultdict(list)
    for row in rows:
        if row["condition"]:
            grouped[row["condition"]].append(row)

    return rows, grouped, condition_status


def fit_size(text: str, base: int, min_size: int = 11, thresholds=None) -> int:
    if thresholds is None:
        thresholds = [(220, -3), (320, -5), (450, -7), (620, -9)]
    length = len(clean(text))
    size = base
    for threshold, delta in thresholds:
        if length > threshold:
            size = base + delta
    return max(min_size, size)


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=18,
    color=C["ink"],
    bold=False,
    font=FONT_BODY,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.04,
    line_spacing=1.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.line_spacing = line_spacing
    for run in p.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def add_multiline(
    slide,
    lines,
    x,
    y,
    w,
    h,
    size=16,
    color=C["ink"],
    bullet=False,
    font=FONT_BODY,
    bold_first=False,
    line_spacing=0.95,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"- {line}" if bullet else line
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        for run in p.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold_first and i == 0
            run.font.color.rgb = rgb(color)
    return box


def add_rect(slide, x, y, w, h, fill, line=None, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1.0)
    else:
        shape.line.fill.background()
    return shape


def add_line(slide, x1, y1, x2, y2, color=C["yellow"], width=3):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    return line


def add_badge(slide, text, x, y, w, fill, color=C["white"], size=11):
    add_rect(slide, x, y, w, 0.28, fill, None, True)
    add_text(slide, text, x + 0.06, y + 0.045, w - 0.12, 0.18, size=size, color=color, bold=True, align=PP_ALIGN.CENTER, font=FONT_HEAD)


def draw_curve(slide):
    points = []
    for i in range(0, 90):
        t = i / 89
        x = -0.2 + 13.9 * t
        y = 6.1 + 1.25 * (1 - (2 * t - 0.35) ** 2)
        points.append((Inches(x), Inches(y)))
    fb = slide.shapes.build_freeform(points[0][0], points[0][1])
    fb.add_line_segments(points[1:], close=False)
    curve = fb.convert_to_shape()
    curve.fill.background()
    curve.line.color.rgb = rgb(C["yellow"])
    curve.line.width = Pt(5.0)
    return curve


def add_header(slide, title, subtitle=None, number=None, section=None):
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, C["white"], None, False)
    draw_curve(slide)
    add_rect(slide, 0.0, 0.0, 0.08, 7.5, C["yellow"], None, False)
    add_text(slide, title, 0.36, 0.22, 8.7, 0.45, size=22, color=C["ink"], bold=True, font=FONT_HEAD)
    if subtitle:
        add_text(slide, subtitle, 0.38, 0.68, 8.8, 0.32, size=11, color=C["muted"], font=FONT_BODY)
    if number is not None:
        add_text(slide, str(number), 11.25, 7.0, 1.55, 0.24, size=10, color=C["muted"], align=PP_ALIGN.RIGHT, font=FONT_HEAD)
    if section:
        add_badge(slide, section, 10.4, 0.28, 2.25, C["taupe"], C["white"], size=9)


def add_stat_card(slide, x, y, w, h, value, label, fill=C["soft_yellow"], accent=C["yellow"]):
    add_rect(slide, x, y, w, h, fill, C["line"], True)
    add_rect(slide, x, y, 0.09, h, accent, None, False)
    add_text(slide, value, x + 0.18, y + 0.22, w - 0.3, 0.45, size=24, color=C["ink"], bold=True, font=FONT_HEAD)
    add_text(slide, label, x + 0.2, y + 0.78, w - 0.35, 0.38, size=11, color=C["muted"])


def add_condition_card(slide, x, y, w, h, item):
    status_fill = C["green"] if item["status"] == "Yes" else C["gray"] if item["status"] == "No" else C["brown"]
    add_rect(slide, x, y, w, h, C["paper"], C["line"], True)
    add_rect(slide, x, y, 0.09, h, status_fill, None, False)
    add_badge(slide, f"Sample: {item['status']}", x + 0.18, y + 0.18, 1.45, status_fill, C["white"], size=8)
    add_text(slide, item["short"], x + 0.2, y + 0.56, w - 0.4, 0.42, size=13, color=C["ink"], bold=True, font=FONT_HEAD)
    add_text(slide, f"{item['active_controls']} active / {item['control_count']} controls", x + 0.2, y + 1.02, w - 0.35, 0.24, size=9.5, color=C["muted"], font=FONT_BODY)


def summarize_group(grouped, condition):
    rows = grouped.get(condition, [])
    return {
        "control_count": len(rows),
        "active_controls": sum(1 for row in rows if row["applicable"]),
        "phases": sorted({row["phase"] for row in rows if row["phase"]}),
        "gates": sorted({row["gate"] for row in rows if row["gate"]}),
        "evidence": sorted({row["evidence"] for row in rows if row["evidence"]}),
        "domains": sorted({row["domain"] for row in rows if row["domain"]}),
        "owners": sorted({row["owner"] for row in rows if row["owner"]}),
    }


def chip_row(slide, title, values, x, y, max_items=4):
    add_text(slide, title, x, y, 1.25, 0.24, size=10, color=C["muted"], bold=True, font=FONT_HEAD)
    cx = x + 1.18
    shown = values[:max_items]
    for value in shown:
        label = value if len(value) <= 34 else value[:31] + "..."
        width = min(2.25, max(0.82, len(label) * 0.07 + 0.34))
        add_badge(slide, label, cx, y - 0.02, width, C["soft"], C["ink"], size=7.5)
        cx += width + 0.12
    if len(values) > max_items:
        add_badge(slide, f"+{len(values) - max_items}", cx, y - 0.02, 0.55, C["soft"], C["ink"], size=7.5)


def add_action_slide(prs, slide_no, item, grouped, condition_status):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    summary = summarize_group(grouped, item["condition"])
    status = condition_status.get(item["condition"], "Not set")
    section = "APPLIES" if status == "Yes" else "CONDITIONAL"
    add_header(slide, item["title"], item["subtitle"], slide_no, section)

    status_fill = C["green"] if status == "Yes" else C["gray"] if status == "No" else C["brown"]
    add_badge(slide, f"Current sample: {status}", 0.42, 1.15, 2.15, status_fill, C["white"], size=9)
    add_badge(slide, f"{summary['active_controls']} active / {summary['control_count']} total controls", 2.75, 1.15, 2.6, C["soft"], C["ink"], size=8.5)

    add_rect(slide, 0.42, 1.58, 5.55, 4.56, C["paper"], C["line"], True)
    add_rect(slide, 0.42, 1.58, 0.1, 4.56, C["yellow"], None, False)
    add_text(slide, "ทีมโปรเจกต์ต้องทำอะไร", 0.68, 1.82, 4.85, 0.32, size=16, color=C["ink"], bold=True, font=FONT_HEAD)
    action_size = item.get("action_size", fit_size(" ".join(item["actions"]), 15, min_size=11))
    add_multiline(slide, item["actions"], 0.72, 2.25, 4.9, 3.62, size=action_size, color=C["ink"], bullet=True, line_spacing=0.9)

    add_rect(slide, 6.22, 1.58, 6.55, 2.18, C["taupe"], None, True)
    add_text(slide, "เล่าแบบง่าย", 6.52, 1.86, 5.75, 0.3, size=15, color=C["yellow"], bold=True, font=FONT_HEAD)
    explain_size = fit_size(item["plain"], 18, min_size=13, thresholds=[(120, -1), (180, -3), (250, -5)])
    add_text(slide, item["plain"], 6.52, 2.27, 5.75, 1.08, size=explain_size, color=C["white"], bold=False, font=FONT_BODY, line_spacing=0.95)

    add_rect(slide, 6.22, 3.98, 6.55, 2.16, C["soft_yellow"], C["line"], True)
    add_text(slide, "Evidence ที่ควรเตรียม", 6.52, 4.2, 5.75, 0.3, size=15, color=C["ink"], bold=True, font=FONT_HEAD)
    evidence = item.get("evidence") or summary["evidence"]
    evidence_lines = evidence[:5]
    add_multiline(slide, evidence_lines, 6.55, 4.58, 5.65, 1.25, size=11.5, color=C["ink"], bullet=True, line_spacing=0.9)

    meta_y = 6.36
    chip_row(slide, "Phase", summary["phases"] or item.get("phase", []), 0.6, meta_y, max_items=3)
    chip_row(slide, "Gate", summary["gates"] or item.get("gate", []), 6.2, meta_y, max_items=2)
    return slide


def build_condition_items(grouped, condition_status):
    def status(condition):
        return condition_status.get(condition, "Not set")

    items = [
        {
            "condition": "ทุกกรณี",
            "short": "ทุกกรณี",
            "title": "ทุกโปรเจกต์ต้องมี baseline controls",
            "subtitle": "ใช้กับทุก AI use case ตั้งแต่เริ่มคิดจนถึง monitoring",
            "plain": "ไม่ว่าโปรเจกต์จะเสี่ยงระดับไหน ทีมต้องเริ่มจากการนิยามความสำเร็จ ลงทะเบียน use case ประเมินความเสี่ยง และเก็บหลักฐานตลอด lifecycle.",
            "actions": [
                "กำหนด AI Model Evaluation Criteria / Business KPI ใน BRD และส่ง AI Risk Assessment ให้ OpRisk review",
                "ลงทะเบียน AI Use Case / Model ใน Model Inventory และเก็บ Registration ID",
                "ทำ explainability section, validation-set test และ Model Performance Report",
                "ระบุ data storage / retention, criticality tier, RPO/RTO และ HA/Backup/DR ใน ARB pack",
                "ตั้ง monitoring dashboard/report, encryption/access control และ rate limit หลัง go-live",
            ],
            "evidence": ["BRD", "AI Risk Assessment Form", "Registration ID", "AI Safety Checklist", "Model Performance Report", "ARB Pack", "Monitoring Dashboard"],
            "status": status("ทุกกรณี"),
            "action_size": 12.5,
        },
        {
            "condition": "มีการนำระบบ AI มาใช้กับงานหลักที่เกี่ยวข้องกับการตัดสินใจเชิงกลยุทธ์(ตามประกาศ BOT) หรือ งานที่ส่งผลกระทบโดยตรงต่อลูกค้า",
            "short": "Strategic / customer-impact AI",
            "title": "ถ้า AI กระทบงานหลักหรือลูกค้า ต้องมี human oversight",
            "subtitle": "แสดงให้เห็นว่าใครตรวจ ใครอนุมัติ และหยุดระบบได้เมื่อเกิดปัญหา",
            "plain": "ถ้า AI มีผลต่อ decision สำคัญหรือ customer impact ต้องออกแบบให้คนยังรับผิดชอบและควบคุมได้จริง ไม่ใช่ปล่อยให้ AI ตัดสินใจลำพัง.",
            "actions": [
                "เลือก oversight model ตาม risk level: HITL หรือ HOTL",
                "ใส่ oversight step ใน customer journey / use case diagram / UI flow",
                "ระบุจุดที่มนุษย์ review, approve, override หรือ stop ระบบได้",
                "แนบภาพ process flow หรือ UI capture เป็นหลักฐาน",
            ],
            "evidence": ["Business process flow", "Customer Journey", "Use Case Diagram", "UI capture"],
            "status": status("มีการนำระบบ AI มาใช้กับงานหลักที่เกี่ยวข้องกับการตัดสินใจเชิงกลยุทธ์(ตามประกาศ BOT) หรือ งานที่ส่งผลกระทบโดยตรงต่อลูกค้า"),
        },
        {
            "condition": "AI Risk Assessment = Medium or High",
            "short": "Risk = Medium or High",
            "title": "ถ้า risk เป็น Medium/High ต้องทดสอบ edge case",
            "subtitle": "อย่าวัดแค่ average performance ต้องดูกรณียากและกรณีผิดพลาดด้วย",
            "plain": "โปรเจกต์ที่มีความเสี่ยงระดับกลางขึ้นไปต้องพิสูจน์ว่าโมเดลรับมือกับกรณีสุดขอบได้ ไม่ใช่ผ่านเฉพาะข้อมูลทั่วไป.",
            "actions": [
                "เตรียม edge case dataset ที่สะท้อนกรณียากหรือกรณีผิดปกติ",
                "ทดสอบผลเทียบกับ criteria ที่กำหนดไว้",
                "รายงานผลใน Model Performance Report",
                "บันทึกข้อจำกัดและ mitigation ถ้าพบ failure pattern",
            ],
            "evidence": ["Edge case dataset", "Model Performance Report"],
            "status": status("AI Risk Assessment = Medium or High"),
        },
        {
            "condition": "AI Risk Assessment = High",
            "short": "Risk = High",
            "title": "ถ้า risk เป็น High ต้องมี monitoring และ kill switch",
            "subtitle": "High-risk AI ต้องมีเกณฑ์เตือน หยุดใช้ และแผนรับเหตุการณ์",
            "plain": "งานเสี่ยงสูงต้องถูก monitor แบบมี threshold ชัดเจน และต้องรู้ล่วงหน้าว่าถ้าคุณภาพตกหรือเกิด incident จะหยุดระบบอย่างไร.",
            "actions": [
                "ทดสอบ robustness ด้วย noisy / perturbed data และบันทึกผล",
                "ตั้ง threshold ของ model quality และ alert เมื่อคุณภาพต่ำกว่าเกณฑ์",
                "ออกแบบ kill switch / disable agent mechanism",
                "ระบุ owner, response path และแผนหยุดใช้ AI เมื่อจำเป็น",
            ],
            "evidence": ["Model Performance Report", "Monitoring Dashboard", "Alert configuration", "Incident Response Plan", "Stop-use plan"],
            "status": status("AI Risk Assessment = High"),
        },
        {
            "condition": "ข้อมูลขาเข้าหรือขาออกจากระบบ AI ถูกจัดเตรียมให้เป็นรายบุคคล",
            "short": "Personalized input/output",
            "title": "ถ้า input/output เป็นรายบุคคล ต้องทดสอบ fairness",
            "subtitle": "ตรวจว่าโมเดลไม่สร้างผลลัพธ์ที่เอนเอียงต่อกลุ่มใดกลุ่มหนึ่ง",
            "plain": "เมื่อผลลัพธ์ AI ถูกปรับตามบุคคล ต้องมีหลักฐานว่าทดสอบ fairness แล้ว โดยเฉพาะกรณีที่เกี่ยวข้องกับ sensitive attributes.",
            "actions": [
                "เลือก fairness metrics ที่เหมาะกับ use case เช่น demographic parity หรือ equalized odds",
                "ทดสอบบน sensitive attributes หรือ proxy ที่เกี่ยวข้อง",
                "บันทึกผลและ limitation ใน AI Safety & Principles Checklist",
                "กำหนด mitigation ถ้าพบ bias หรือ disparate impact",
            ],
            "evidence": ["Fairness test result", "AI Safety & Principles Checklist"],
            "status": status("ข้อมูลขาเข้าหรือขาออกจากระบบ AI ถูกจัดเตรียมให้เป็นรายบุคคล"),
        },
        {
            "condition": "Business-critical supervised model",
            "short": "Business-critical supervised model",
            "title": "ถ้าเป็น supervised model ในงานสำคัญ ต้องอธิบายผลลัพธ์ได้",
            "subtitle": "ให้คนตรวจสอบได้ว่าโมเดลใช้เหตุผลหรือ feature อะไรประกอบการตัดสินใจ",
            "plain": "สำหรับโมเดลที่ช่วยตัดสินใจเรื่องธุรกิจสำคัญ ทีมต้องมี explainability ที่อ่านได้ ไม่ใช่มีแค่คะแนน prediction.",
            "actions": [
                "ใช้ SHAP, LIME หรือ feature importance ตามความเหมาะสม",
                "อธิบายปัจจัยสำคัญที่มีผลต่อ output ของโมเดล",
                "ใส่ explainability section ใน Model Performance Report",
                "ชี้ limitation ว่าคำอธิบายใช้ได้ในขอบเขตใด",
            ],
            "evidence": ["Model Performance Report", "Explainability section"],
            "status": status("Business-critical supervised model"),
        },
        {
            "condition": "เป็น Generative AI",
            "short": "Generative AI",
            "title": "ถ้าเป็น Generative AI ต้องอ้างอิงแหล่งข้อมูลได้",
            "subtitle": "ลด hallucination ด้วย RAG, source citation หรือหลักฐานการอ้างอิง",
            "plain": "ถ้า AI สร้างคำตอบหรือเนื้อหาเอง ต้องมีวิธีให้คนตรวจที่มาของคำตอบได้ โดยเฉพาะคำตอบที่มีผลต่อธุรกิจหรือลูกค้า.",
            "actions": [
                "ออกแบบ RAG หรือ source citation ใน prompt / response flow",
                "เก็บ screenshot ที่เห็น reference หรือ reasoning trace ที่ตรวจสอบได้",
                "ระบุข้อจำกัดของ citation และกรณีที่ต้องให้คน review",
                "จัดเก็บตัวอย่าง output ที่ผ่านและไม่ผ่านไว้ใช้ตรวจซ้ำ",
            ],
            "evidence": ["Screen capture", "Prompt / response example", "Citation evidence"],
            "status": status("เป็น Generative AI"),
        },
        {
            "condition": "ใช้งาน External data ซึ่งเป็น Open data",
            "short": "External open data",
            "title": "ถ้าใช้ open data ต้องรู้แหล่งที่มาและ license",
            "subtitle": "Open data ไม่ได้แปลว่าใช้ได้ทุกกรณี ต้องมี provenance",
            "plain": "ทีมต้องบอกได้ว่าข้อมูลมาจากไหน ใครรวบรวม และ license อนุญาตให้ใช้กับ use case นี้หรือไม่.",
            "actions": [
                "ระบุ data provenance: source, collector, version/date และ license",
                "ใส่ข้อมูล provenance ใน ARB Pack",
                "ตรวจว่าการใช้งานสอดคล้องกับ license และวัตถุประสงค์ของโปรเจกต์",
                "เก็บ snapshot หรือ link ของ source ที่ใช้จริง",
            ],
            "evidence": ["ARB Pack", "Data provenance record", "License evidence"],
            "status": status("ใช้งาน External data ซึ่งเป็น Open data"),
        },
        {
            "condition": "ใช้สนทนาหรือสื่อสารกับลูกค้าแทนมนุษย์",
            "short": "Customer-facing conversation AI",
            "title": "ถ้า AI คุยกับลูกค้า ต้องเปิดเผยให้ชัด",
            "subtitle": "ลูกค้าต้องรู้ว่ากำลังคุยกับ AI และข้อมูลถูกใช้เพื่ออะไร",
            "plain": "AI ที่สื่อสารกับลูกค้าไม่ควรถูกทำให้ดูเหมือนมนุษย์โดยไม่แจ้ง ต้องมี disclosure และหลักฐานหน้าจอก่อน go-live.",
            "actions": [
                "เพิ่มข้อความ disclosure บน UI / chatbot ว่ากำลังให้บริการด้วย AI",
                "แจ้งวัตถุประสงค์การใช้ข้อมูลให้ลูกค้าเข้าใจ",
                "ตรวจ wording ให้ชัด อ่านง่าย และอยู่ในจุดที่เห็นก่อนใช้บริการ",
                "capture screenshot เก็บเป็น evidence",
            ],
            "evidence": ["Screen capture", "UI / chatbot disclosure wording"],
            "status": status("ใช้สนทนาหรือสื่อสารกับลูกค้าแทนมนุษย์"),
        },
        {
            "condition": "ทำ Model Training หรือ Finetuning",
            "short": "Model training / fine-tuning",
            "title": "ถ้ามี training หรือ fine-tuning ต้อง trace ทุกอย่างได้",
            "subtitle": "ข้อมูล โมเดล พารามิเตอร์ และผลทดลองต้องย้อนกลับได้",
            "plain": "การ train หรือ fine-tune ทำให้ทีมต้องรับผิดชอบ lineage มากขึ้น ต้องรู้ว่าข้อมูลไหน โมเดลไหน และ run ไหนทำให้เกิดผลลัพธ์.",
            "actions": [
                "ใช้ Model Registry / Experiment Tracking บันทึก base model, dataset version, hyperparameters และ run id",
                "นิยาม data quality criteria และให้ Data Owner sign-off ก่อน train/fine-tune",
                "ทำ Data Profiling / EDA report และแนบใน ARB Pack",
                "ทำ Model Card, input/output schema, data lineage และ model lineage",
                "ออกแบบ feedback collection พร้อม retention และ usage ใน ARB Pack",
            ],
            "evidence": ["Model Training Report", "Data quality criteria", "EDA Report", "Model Card", "Lineage record", "ARB Pack"],
            "status": status("ทำ Model Training หรือ Finetuning"),
            "action_size": 11.7,
        },
        {
            "condition": "Deploy ระบบ AI ภายในระบบธนาคาร",
            "short": "Internal bank deployment",
            "title": "ถ้า deploy ในระบบธนาคาร ต้องพร้อม audit และ rollback",
            "subtitle": "ทุก version, deploy log, security control และ fallback path ต้องตรวจย้อนหลังได้",
            "plain": "Production AI ต้องบริหารเหมือนระบบสำคัญ: รู้ว่า deploy อะไร เมื่อไร ใครทำ และถ้าพังจะ rollback หรือส่งงานกลับให้คนอย่างไร.",
            "actions": [
                "ใช้ semantic versioning กับ model ที่ deploy และบันทึกใน Model Deployment Record",
                "เก็บ deploy log: version, timestamp, deployer และ CI/CD evidence",
                "ออกแบบ rollback procedure และ human fallback path",
                "ทำ encryption at rest/in transit, RBAC และ access log สำหรับ artifact และ endpoint",
            ],
            "evidence": ["Model Deployment Record", "Deploy log", "Incident Response Plan", "Security control evidence"],
            "status": status("Deploy ระบบ AI ภายในระบบธนาคาร"),
        },
        {
            "condition": "ใช้ Model จาก Vendor / 3rd Party",
            "short": "Vendor / 3rd-party model",
            "title": "ถ้าใช้ vendor model ต้องคุมสัญญา SLA และ risk assessment",
            "subtitle": "ซื้อหรือใช้ model คนอื่นไม่ได้ลดความรับผิดชอบของโปรเจกต์",
            "plain": "Vendor model ต้องมีหลักฐานว่าเลือกอย่างรอบคอบ มีเงื่อนไขสัญญาคุ้มครองข้อมูล มี SLA และตรวจ model risk ก่อน onboarding.",
            "actions": [
                "ระบุ data boundary, no-internal-data-training และ liability clauses ในสัญญา",
                "เก็บ License / ToS และส่ง Legal review ก่อนใช้งาน",
                "กำหนด SLA: uptime, response time, support และรับ SLA report ตามรอบ",
                "ทำ Vendor Due Diligence / Vendor Risk Assessment ตาม Procurement Process",
                "ตรวจ vendor security และทำ Model Risk Assessment ใน ARB Pack",
            ],
            "evidence": ["Vendor Management Due Diligence", "Legal review", "Contract clauses", "Vendor SLA report", "ARB Pack", "Monitoring Dashboard"],
            "status": status("ใช้ Model จาก Vendor / 3rd Party"),
            "action_size": 12,
        },
        {
            "condition": "ใช้งาน External data ซึ่งได้รับมาจาก vendor / 3rd party",
            "short": "External data from vendor / 3rd party",
            "title": "ถ้าใช้ข้อมูลจาก vendor ต้องผ่าน procurement, legal และ PDPA",
            "subtitle": "สิทธิการใช้ข้อมูลต้องถูกตรวจตั้งแต่ก่อนนำเข้า use case",
            "plain": "ข้อมูลจาก third party ต้องมีสิทธิใช้ชัดเจน และ contract ต้องบอกขอบเขตการใช้ การเก็บ และการปกป้องข้อมูล.",
            "actions": [
                "ดำเนินการตาม Procurement Process",
                "ให้ Legal / PDPA team ตรวจ contract เรื่อง data จาก 3rd party",
                "ระบุสิทธิการใช้, retention, sharing restriction และ security obligation",
                "เก็บ Data Governance Checklist เป็นหลักฐาน",
            ],
            "evidence": ["Data Governance Checklist", "Contract / Legal review", "PDPA review"],
            "status": status("ใช้งาน External data ซึ่งได้รับมาจาก vendor / 3rd party"),
        },
        {
            "condition": "ใช้งาน Generative AI / LLM / prompt‑based system",
            "short": "LLM / prompt-based system",
            "title": "ถ้าใช้ LLM หรือ prompt-based system ต้องทดสอบ prompt attack",
            "subtitle": "ก่อน go-live ต้องรู้ว่าระบบทน prompt injection / jailbreak แค่ไหน",
            "plain": "ระบบที่รับ prompt มีช่องทางถูกสั่งให้ละเมิด policy หรือเปิดเผยข้อมูล ต้องทดสอบก่อนใช้งานจริงและเก็บหลักฐานใน security review.",
            "actions": [
                "อ้างอิง AI Security Guideline",
                "ทดสอบ prompt injection และ jailbreak ก่อน go-live",
                "บันทึก test case, result และ mitigation ที่ใช้",
                "แนบหลักฐานใน System Architecture / Security review",
            ],
            "evidence": ["System Architecture Design", "Prompt injection test result", "Jailbreak test result"],
            "status": status("ใช้งาน Generative AI / LLM / prompt‑based system"),
        },
        {
            "condition": "ใช้ Open Source Model",
            "short": "Open-source model",
            "title": "ถ้าใช้ open-source model ต้องตรวจ license compatibility",
            "subtitle": "เปิด source ไม่ได้แปลว่าใช้ได้อิสระในทุก business context",
            "plain": "ทีมต้องยืนยันว่า license ของ model ใช้กับวัตถุประสงค์นี้ได้ และไม่สร้างข้อผูกพันที่กระทบระบบหรือข้อมูลขององค์กร.",
            "actions": [
                "อ้างอิง DevSecOps Open Source Guideline",
                "เก็บหลักฐาน license ของ model และ dependency สำคัญ",
                "ตรวจ license compatibility กับ use case และ deployment model",
                "แนบผลตรวจใน architecture / security evidence",
            ],
            "evidence": ["System Architecture Design", "Open-source license evidence", "License compatibility review"],
            "status": status("ใช้ Open Source Model"),
        },
        {
            "condition": "ใช้งาน External data ซึ่งเป็น Non-Open data เช่น web scraping",
            "short": "Non-open external data / web scraping",
            "title": "ถ้าใช้ non-open external data ต้องตรวจ PDPA, copyright และ cyber risk",
            "subtitle": "โดยเฉพาะข้อมูลจาก web scraping หรือแหล่งที่ไม่ได้เปิดให้ใช้ชัดเจน",
            "plain": "ข้อมูลภายนอกที่ไม่ใช่ open data ต้องถูกตรวจสิทธิและความเสี่ยงก่อนนำมาใช้ เพราะอาจกระทบ PDPA, copyright หรือ cyber crime.",
            "actions": [
                "อ้างอิง PDPA for AI และ guideline ที่เกี่ยวข้อง",
                "ตรวจ copyright และ cyber crime risk ของแหล่งข้อมูลภายนอก",
                "บันทึกแหล่งข้อมูล วิธีได้มา และเหตุผลที่ใช้ได้",
                "เก็บผลตรวจใน Data Governance Checklist",
            ],
            "evidence": ["Data Governance Checklist", "PDPA / Copyright review", "Source risk record"],
            "status": status("ใช้งาน External data ซึ่งเป็น Non-Open data เช่น web scraping"),
        },
    ]
    return items


def add_cover(prs, rows, grouped, condition_status):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, C["white"], None, False)
    draw_curve(slide)
    add_rect(slide, 0.72, 0.0, 1.28, 1.55, C["taupe"], None, False)
    add_rect(slide, 0.9, 0.28, 0.92, 0.38, C["yellow"], None, True)
    add_text(slide, "AI-SLC", 0.77, 0.82, 1.18, 0.26, size=14, color=C["white"], bold=True, align=PP_ALIGN.CENTER, font=FONT_HEAD)
    add_text(slide, "Action Guide", 0.79, 1.1, 1.14, 0.22, size=8.5, color=C["white"], align=PP_ALIGN.CENTER, font=FONT_BODY)
    add_text(slide, "Applied Condition\nAction Guide", 1.6, 2.9, 7.8, 1.12, size=36, color=C["ink"], bold=True, font=FONT_HEAD, line_spacing=0.9)
    add_text(slide, "อธิบายว่าแต่ละ Applied condition for project ต้องทำอะไร เตรียม evidence อะไร และเกี่ยวกับ phase/gate ไหน", 1.62, 4.18, 8.5, 0.48, size=15, color=C["muted"], font=FONT_BODY)
    add_stat_card(slide, 1.62, 5.25, 2.25, 1.15, str(len(grouped)), "condition groups")
    add_stat_card(slide, 4.12, 5.25, 2.25, 1.15, str(len(rows)), "control rows")
    add_stat_card(slide, 6.62, 5.25, 2.25, 1.15, str(sum(1 for r in rows if r["applicable"])), "active in sample")
    add_text(slide, "Style reference: AISLC PowerPoint\nSource: Control_Guideline sheet", 10.15, 6.45, 2.55, 0.46, size=9.5, color=C["muted"], align=PP_ALIGN.RIGHT)
    return slide


def add_how_to_read(prs, slide_no):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "อ่าน deck นี้แบบเร็ว", "ใช้เป็น script สำหรับเล่าให้ project team เข้าใจว่า condition แต่ละแบบแปลว่าอะไร", slide_no, "HOW TO USE")
    cards = [
        ("1", "ดู condition", "ถ้าโปรเจกต์ตอบ Yes ใน condition นั้น ให้ถือว่า control กลุ่มนี้ถูกเปิดใช้"),
        ("2", "ทำ action", "ทำเฉพาะงานที่จำเป็นต่อ condition นั้น ไม่ต้องอ่าน checklist ทั้ง sheet ทีละแถว"),
        ("3", "เก็บ evidence", "ทุก action ต้องมี artifact ให้ review ได้ เช่น BRD, ARB Pack, report หรือ screenshot"),
        ("4", "เช็ค phase/gate", "รู้ว่าต้องเตรียมหลักฐานก่อนเข้า phase หรือ gate ไหนของ AI-SLC"),
    ]
    x_positions = [0.6, 3.78, 6.96, 10.14]
    for x, (num, title, body) in zip(x_positions, cards):
        add_rect(slide, x, 1.55, 2.55, 3.65, C["paper"], C["line"], True)
        add_rect(slide, x + 0.22, 1.85, 0.52, 0.52, C["yellow"], None, True)
        add_text(slide, num, x + 0.31, 1.94, 0.34, 0.22, size=13, color=C["ink"], bold=True, align=PP_ALIGN.CENTER, font=FONT_HEAD)
        add_text(slide, title, x + 0.22, 2.62, 2.08, 0.3, size=15, color=C["ink"], bold=True, font=FONT_HEAD)
        add_text(slide, body, x + 0.22, 3.08, 2.08, 1.35, size=12, color=C["muted"], font=FONT_BODY, line_spacing=0.92)
    add_rect(slide, 1.0, 5.72, 11.1, 0.74, C["taupe"], None, True)
    add_text(slide, "คำอธิบายใน deck นี้เป็นเวอร์ชันเล่าง่าย แต่ยังอิง action/evidence จาก Control_Guideline", 1.28, 5.95, 10.55, 0.26, size=13.5, color=C["white"], align=PP_ALIGN.CENTER, font=FONT_BODY)
    return slide


def add_sample_summary(prs, slide_no, condition_items):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "โปรเจกต์ตัวอย่างเปิด control กลุ่มไหนบ้าง", "อิงค่า Yes/No จาก sheet Applied condition for project และ Applicable ใน Control_Guideline", slide_no, "SAMPLE")
    yes_items = [item for item in condition_items if item["status"] == "Yes"]
    no_items = [item for item in condition_items if item["status"] != "Yes"]
    add_text(slide, "Applied = Yes", 0.62, 1.22, 3.6, 0.38, size=18, color=C["ink"], bold=True, font=FONT_HEAD)
    for i, item in enumerate(yes_items[:8]):
        add_condition_card(slide, 0.62 + (i % 2) * 3.2, 1.72 + (i // 2) * 1.35, 2.85, 1.08, item)
    add_text(slide, "Conditional / Not active in sample", 7.2, 1.22, 4.5, 0.38, size=18, color=C["ink"], bold=True, font=FONT_HEAD)
    for i, item in enumerate(no_items[:8]):
        add_condition_card(slide, 7.2 + (i % 2) * 2.78, 1.72 + (i // 2) * 1.18, 2.45, 0.93, item)
    add_text(slide, "หมายเหตุ: สไลด์ถัดไปยังอธิบายกลุ่มที่เป็น No ไว้ด้วย เพื่อใช้ generic กับโปรเจกต์อื่นที่อาจตอบ Yes", 7.2, 6.64, 5.1, 0.34, size=10.5, color=C["muted"])
    return slide


def add_close(prs, slide_no):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, "วิธีใช้เป็น operating checklist", "สรุป flow ที่ project team ควรทำก่อนส่ง review", slide_no, "WRAP")
    steps = [
        ("1", "ตอบ condition", "เลือก Yes/No ให้ครบใน Applied condition for project"),
        ("2", "ดึง control", "ดู Control_Guideline เฉพาะแถว Applicable = TRUE"),
        ("3", "ทำ evidence pack", "รวม BRD, ARB Pack, report, checklist, screenshot และ approval"),
        ("4", "ส่งตาม gate", "ส่ง evidence ให้ owner/gate controller ก่อนเดิน phase ถัดไป"),
    ]
    for i, (num, title, body) in enumerate(steps):
        x = 0.92 + i * 3.05
        add_rect(slide, x, 1.72, 2.45, 3.5, C["paper"], C["line"], True)
        add_rect(slide, x + 0.82, 1.42, 0.82, 0.82, C["yellow"], None, True)
        add_text(slide, num, x + 1.02, 1.59, 0.42, 0.3, size=18, color=C["ink"], bold=True, align=PP_ALIGN.CENTER, font=FONT_HEAD)
        add_text(slide, title, x + 0.22, 2.55, 2.0, 0.32, size=15, color=C["ink"], bold=True, align=PP_ALIGN.CENTER, font=FONT_HEAD)
        add_text(slide, body, x + 0.22, 3.08, 2.0, 1.0, size=12.2, color=C["muted"], align=PP_ALIGN.CENTER, font=FONT_BODY, line_spacing=0.92)
        if i < len(steps) - 1:
            add_line(slide, x + 2.55, 3.2, x + 2.95, 3.2, C["yellow"], 4)
    add_rect(slide, 1.22, 5.82, 10.9, 0.7, C["taupe"], None, True)
    add_text(slide, "หลักจำง่าย: condition เปิด control, control เปิด action, action ต้องจบด้วย evidence", 1.44, 6.04, 10.44, 0.25, size=14, color=C["white"], align=PP_ALIGN.CENTER, font=FONT_HEAD)
    return slide


def set_doc_properties(prs):
    props = prs.core_properties
    props.title = "Applied Condition Action Guide"
    props.subject = "AI-SLC control guideline summary"
    props.author = "Codex"
    props.keywords = "AI-SLC, applied condition, control guideline, AI risk assessment"
    props.comments = "Generated from Control_Guideline sheet and styled after AISLC reference deck."


def write_outline(condition_items, rows, grouped, condition_status):
    lines = [
        "# AISLC Applied Condition Action Guide",
        "",
        "Audience: project teams and stakeholders who need to understand which AI-SLC controls apply to a project.",
        "Objective: explain each `Applied condition for project` in simple terms, with required actions, evidence, and phase/gate context.",
        "Rendering mode: editable PowerPoint, no image generation.",
        "Aspect ratio: 16:9.",
        "",
        "## Source Summary",
        f"- Source workbook: `{SOURCE_XLSX}`",
        f"- Source style deck: `{SOURCE_PPTX}`",
        f"- Control rows used: {len(rows)}",
        f"- Condition groups used: {len(grouped)}",
        f"- Active controls in sample: {sum(1 for row in rows if row['applicable'])}",
        "",
        "## Slide Outline",
        "1. Applied Condition Action Guide - cover",
        "2. อ่าน deck นี้แบบเร็ว - how to use",
        "3. โปรเจกต์ตัวอย่างเปิด control กลุ่มไหนบ้าง - sample summary",
    ]
    for idx, item in enumerate(condition_items, start=4):
        summary = summarize_group(grouped, item["condition"])
        lines.append(f"{idx}. {item['title']} - {item['short']} ({item['status']}; {summary['active_controls']}/{summary['control_count']} active controls)")
    lines.append(f"{len(condition_items) + 4}. วิธีใช้เป็น operating checklist - wrap-up")
    lines.extend(
        [
            "",
            "## Applied Condition Status",
        ]
    )
    for item in condition_items:
        summary = summarize_group(grouped, item["condition"])
        lines.append(f"- {item['short']}: sample={item['status']}; controls={summary['control_count']}; active={summary['active_controls']}")
    OUTLINE.write_text("\n".join(lines) + "\n", encoding="utf-8")

    plan = {
        "project": PROJECT,
        "source_workbook": str(SOURCE_XLSX),
        "source_style_deck": str(SOURCE_PPTX),
        "output_pptx": str(OUT_PPTX),
        "slides": [
            {"title": "Applied Condition Action Guide", "role": "cover"},
            {"title": "อ่าน deck นี้แบบเร็ว", "role": "how-to-use"},
            {"title": "โปรเจกต์ตัวอย่างเปิด control กลุ่มไหนบ้าง", "role": "sample-summary"},
        ]
        + [
            {
                "title": item["title"],
                "condition": item["condition"],
                "short": item["short"],
                "sample_status": item["status"],
                "summary": summarize_group(grouped, item["condition"]),
            }
            for item in condition_items
        ]
        + [{"title": "วิธีใช้เป็น operating checklist", "role": "wrap-up"}],
    }
    PLAN_JSON.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")


def maybe_copy_theme(prs):
    # The user asked for the visual style of the attached PowerPoint, but this deck
    # is generated as editable shapes/text to keep Thai control wording accurate.
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)


def build_deck():
    rows, grouped, condition_status = read_workbook()
    condition_items = build_condition_items(grouped, condition_status)
    for item in condition_items:
        summary = summarize_group(grouped, item["condition"])
        item["control_count"] = summary["control_count"]
        item["active_controls"] = summary["active_controls"]
        if "status" not in item:
            item["status"] = condition_status.get(item["condition"], "Not set")

    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    OUTLINE.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    maybe_copy_theme(prs)
    set_doc_properties(prs)

    add_cover(prs, rows, grouped, condition_status)
    add_how_to_read(prs, 2)
    add_sample_summary(prs, 3, condition_items)
    slide_no = 4
    for item in condition_items:
        add_action_slide(prs, slide_no, item, grouped, condition_status)
        slide_no += 1
    add_close(prs, slide_no)

    prs.save(OUT_PPTX)
    write_outline(condition_items, rows, grouped, condition_status)
    return OUT_PPTX


if __name__ == "__main__":
    out = build_deck()
    print(out)
