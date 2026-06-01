from pathlib import Path
import textwrap

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
PROJECT = "hedge-fund-strategies-research"
IMG_DIR = ROOT / "assets" / "images" / PROJECT
OUT_DIR = ROOT / "outputs"
NOTES_DIR = ROOT / "notes"
IMG_DIR.mkdir(parents=True, exist_ok=True)
OUT_DIR.mkdir(parents=True, exist_ok=True)
NOTES_DIR.mkdir(parents=True, exist_ok=True)

W, H = 1920, 1080
BG = "#0B1117"
PANEL = "#121B24"
PANEL2 = "#182330"
TEXT = "#F4F7FA"
MUTED = "#A7B3C2"
GRID = "#314255"
ACCENT = "#4CC9F0"
GREEN = "#50C878"
AMBER = "#F6C85F"
RED = "#F26D6D"
PURPLE = "#B48CFF"

FONT_REG = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"
FONT_BOLD = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"


def font(size):
    return ImageFont.truetype(FONT_REG, size)


def draw_text(draw, xy, text, size=34, fill=TEXT, max_width=None, line_spacing=1.18):
    f = font(size)
    x, y = xy
    if max_width is None:
        draw.text((x, y), text, font=f, fill=fill)
        return y + size
    lines = []
    for para in str(text).split("\n"):
        words = para.split(" ")
        line = ""
        for word in words:
            test = (line + " " + word).strip()
            if draw.textbbox((0, 0), test, font=f)[2] <= max_width or not line:
                line = test
            else:
                lines.append(line)
                line = word
        lines.append(line)
    for line in lines:
        draw.text((x, y), line, font=f, fill=fill)
        y += int(size * line_spacing)
    return y


def base(title, kicker=None):
    img = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, W, 82), fill="#071018")
    d.line((80, 82, W - 80, 82), fill="#244056", width=2)
    if kicker:
        draw_text(d, (82, 28), kicker.upper(), 22, ACCENT)
    draw_text(d, (80, 118), title, 54, TEXT, 1480, 1.05)
    return img, d


def save_slide(img, idx):
    path = IMG_DIR / f"slide-{idx:02d}.png"
    img.save(path, quality=95)
    return path


def pill(d, x, y, text, color=ACCENT, w=None):
    f = font(24)
    tw = d.textbbox((0, 0), text, font=f)[2]
    ww = w or tw + 36
    d.rounded_rectangle((x, y, x + ww, y + 46), radius=16, fill=color)
    d.text((x + 18, y + 10), text, font=f, fill="#061018")


def card(d, xy, wh, title, body, accent=ACCENT):
    x, y = xy
    w, h = wh
    d.rounded_rectangle((x, y, x + w, y + h), radius=18, fill=PANEL)
    d.rectangle((x, y, x + 8, y + h), fill=accent)
    draw_text(d, (x + 28, y + 24), title, 34, TEXT, w - 60)
    draw_text(d, (x + 28, y + 82), body, 24, MUTED, w - 60, 1.22)


major_returns = pd.DataFrame({
    "Strategy": ["Multi-Strategy", "Long biased", "Event", "Equity L/S", "Credit", "Macro", "Quant", "Arbitrage", "HF Composite"],
    "2025 H1": [3.99, 9.61, 5.51, 4.40, 3.20, 4.80, 1.93, 3.58, 4.50],
    "2024": [13.11, 12.45, 10.19, 13.48, 9.77, 10.05, 8.62, 6.00, 11.47],
    "2023": [7.32, 14.47, 9.51, 11.54, 8.80, 5.78, 1.60, 2.39, 8.86],
    "2022": [9.11, -12.94, -4.67, -9.63, -2.64, 5.82, 8.71, 2.71, -2.67],
    "2021": [11.24, 12.55, 11.97, 5.48, 9.40, 0.29, 7.73, 3.98, 8.27],
    "5Y CAR": [11.08, 10.04, 9.27, 8.10, 7.42, 7.03, 6.41, 5.33, 8.36],
    "5Y Vol": [2.69, 10.77, 5.32, 7.26, 2.90, 3.30, 5.27, 2.34, 4.31],
})

sub_5y = pd.DataFrame([
    ("Event - Activist", 14.4), ("Long - Equity", 13.3), ("Quant - Multi", 12.9),
    ("Long - Commodities", 12.3), ("Multi-strategy", 11.1), ("Arbitrage - Opportunistic", 10.9),
    ("Macro - Commodities", 10.4), ("Equity L/S - Global", 10.2),
    ("Credit - Structured", 8.6), ("Credit - Distressed", 8.5), ("Quant - Risk Premia", 8.5),
    ("Equity L/S - US", 8.4), ("Event - Multi", 8.3), ("Equity L/S - Fundamental EMN", 8.3),
    ("Long - Other", 8.1), ("Equity L/S - Other", 8.0), ("Convertible Arb", 7.8),
    ("Macro - FI Relative Value", 7.6), ("Event - M&A", 7.5), ("Equity L/S - APAC", 7.4),
    ("Quant - EMN", 7.3), ("Event - Opportunistic", 7.2), ("Credit - Direct Lending", 6.9),
    ("Macro - Global", 6.7), ("Macro - EM", 6.7), ("Equity L/S - Sector", 6.3),
    ("Credit - Relative Value", 5.7), ("Credit - Structured LO", 5.3), ("Quant - Macro", 5.3),
    ("Credit - Municipal", 4.2), ("Long - Diversified Growth", 3.4),
], columns=["Sub-strategy", "5Y CAR"])

strategy_examples = [
    ("Multi-strategy / pod shop", "Citadel, Millennium, Point72, Balyasny, Schonfeld, ExodusPoint"),
    ("Global macro / discretionary macro", "Bridgewater Pure Alpha, Brevan Howard, Tudor, Caxton, Rokos"),
    ("Systematic macro / CTA / managed futures", "Man AHL, Winton, Aspect, AQR Managed Futures, Lynx"),
    ("Equity long/short", "Tiger Global, Lone Pine, Coatue, Viking, Point72 sector pods"),
    ("Equity market neutral / stat arb", "D.E. Shaw, Two Sigma, Renaissance, AQR, Citadel GQS"),
    ("Event-driven / merger arbitrage", "Elliott, Paulson, Farallon, Pentwater, Tiedemann/TIG"),
    ("Activist / special situations", "Elliott, Pershing Square, Trian, Third Point, Starboard"),
    ("Credit / distressed / structured", "Oaktree, Baupost, Angelo Gordon/TPG Angelo Gordon, Magnetar, King Street"),
    ("Relative value / convertibles / volatility", "Citadel, Millennium, D.E. Shaw, Balyasny, Capula"),
    ("Commodity / energy", "Citadel Commodities, Millennium, Andurand, Mercuria-linked teams, Man AHL"),
    ("Crypto / digital assets", "Brevan Howard Digital, Pantera, Multicoin, Galaxy, BlockTower"),
]

sources = [
    ("HFR 2025 Global Hedge Fund Industry Report", "https://hfr-wp-s3.s3.amazonaws.com/wp-content/uploads/2026/01/22133429/2025.Q4-HFR-GIR_FINAL.pdf"),
    ("HFRI December 2024 Performance Notes", "https://hfr-wp-s3.s3.amazonaws.com/wp-content/uploads/2025/01/08151305/2024.12_HFRI-Flash.pdf"),
    ("HFRI December 2023 Performance Notes", "https://www.hfr.com/media/performance-notes/hfri-indices-december-2023-performance-notes/"),
    ("HFRI December 2022 Performance Notes", "https://www.hfr.com/media/performance-notes/hfri-indices-december-2022-performance-notes/"),
    ("HFRI December 2021 Performance Notes", "https://www.hfr.com/media/performance-notes/hfri-indices-december-2021-performance-notes/"),
    ("Aurum Hedge Fund Industry Deep Dive H1 2025", "https://www.aurum.com/wp-content/uploads/Aurum-Industry-Deep-Dive-H1-2025-review.pdf"),
    ("Aurum Hedge Fund Industry Deep Dive Chartbook 2025", "https://www.aurum.com/wp-content/uploads/Aurum-Hedge-Fund-Industry-Deep-Dive-Chartbook-2025.pdf"),
    ("Citadel What We Do", "https://www.citadel.com/what-we-do/"),
    ("Millennium Approach", "https://www.mlp.com/approach/"),
    ("Bridgewater All Weather Strategy", "https://www.bridgewater.com/research-and-insights/the-all-weather-strategy"),
    ("Elliott Strategies", "https://www.elliottimuka.com/strategies.html"),
    ("Man AHL", "https://www.man.com/ahl"),
    ("AQR Strategies", "https://www.aqr.com/strategies"),
    ("Two Sigma Investment Management", "https://www.twosigma.com/businesses/investment-management/"),
]


def chart_major_heatmap(path):
    data = major_returns.set_index("Strategy")[["2021", "2022", "2023", "2024", "2025 H1", "5Y CAR"]]
    fig, ax = plt.subplots(figsize=(11.8, 6.0), dpi=160)
    fig.patch.set_facecolor(BG)
    ax.set_facecolor(BG)
    vals = data.values
    im = ax.imshow(vals, cmap="RdYlGn", vmin=-13, vmax=14)
    ax.set_xticks(range(data.shape[1]), data.columns, color=TEXT, fontsize=10)
    ax.set_yticks(range(data.shape[0]), data.index, color=TEXT, fontsize=10)
    for i in range(data.shape[0]):
        for j in range(data.shape[1]):
            ax.text(j, i, f"{vals[i, j]:.1f}%", ha="center", va="center", color="#061018", fontsize=9, fontweight="bold")
    for spine in ax.spines.values():
        spine.set_visible(False)
    ax.tick_params(length=0)
    plt.tight_layout()
    fig.savefig(path, facecolor=BG, bbox_inches="tight")
    plt.close(fig)


def chart_risk_return(path):
    df = major_returns[major_returns["Strategy"] != "HF Composite"]
    fig, ax = plt.subplots(figsize=(8.8, 5.4), dpi=160)
    fig.patch.set_facecolor(BG)
    ax.set_facecolor(BG)
    colors = [ACCENT, GREEN, AMBER, PURPLE, "#7DD3FC", "#F472B6", "#A3E635", "#F97316"]
    ax.scatter(df["5Y Vol"], df["5Y CAR"], s=170, c=colors, edgecolors="white", linewidth=1.2)
    for _, r in df.iterrows():
        ax.text(r["5Y Vol"] + 0.12, r["5Y CAR"], r["Strategy"], color=TEXT, fontsize=9, va="center")
    ax.set_xlabel("5Y Volatility (%)", color=MUTED)
    ax.set_ylabel("5Y CAR (%)", color=MUTED)
    ax.grid(color=GRID, alpha=.65)
    ax.tick_params(colors=MUTED)
    for spine in ax.spines.values():
        spine.set_color(GRID)
    plt.tight_layout()
    fig.savefig(path, facecolor=BG, bbox_inches="tight")
    plt.close(fig)


def chart_sub5(path):
    df = sub_5y.head(18).iloc[::-1]
    fig, ax = plt.subplots(figsize=(10.0, 6.4), dpi=160)
    fig.patch.set_facecolor(BG)
    ax.set_facecolor(BG)
    bars = ax.barh(df["Sub-strategy"], df["5Y CAR"], color=ACCENT)
    for b, v in zip(bars, df["5Y CAR"]):
        ax.text(v + .15, b.get_y() + b.get_height()/2, f"{v:.1f}%", va="center", color=TEXT, fontsize=9)
    ax.set_xlim(0, 16)
    ax.set_xlabel("5Y compound annual return to Jun 2025 (%)", color=MUTED)
    ax.tick_params(colors=MUTED, labelsize=9)
    ax.grid(axis="x", color=GRID, alpha=.65)
    for spine in ax.spines.values():
        spine.set_visible(False)
    plt.tight_layout()
    fig.savefig(path, facecolor=BG, bbox_inches="tight")
    plt.close(fig)


def table_slide(d, rows, x, y, colw, rowh=54, header=None, size=23):
    if header:
        d.rounded_rectangle((x, y, x + sum(colw), y + rowh), radius=12, fill="#203142")
        cx = x
        for i, h in enumerate(header):
            draw_text(d, (cx + 12, y + 14), h, size, TEXT, colw[i] - 18)
            cx += colw[i]
        y += rowh + 6
    for idx, row in enumerate(rows):
        fill = PANEL if idx % 2 == 0 else "#101821"
        d.rounded_rectangle((x, y, x + sum(colw), y + rowh), radius=10, fill=fill)
        cx = x
        for i, cell in enumerate(row):
            draw_text(d, (cx + 12, y + 13), cell, size, TEXT if i == 0 else MUTED, colw[i] - 20, 1.05)
            cx += colw[i]
        y += rowh + 5


def make_slides():
    paths = []

    img, d = base("Hedge Fund Strategies: Research Map & 5-Year Performance", "Research deck | 2021-2025")
    draw_text(d, (82, 305), "กลยุทธ์การลงทุนของ hedge fund ไม่ได้มีแบบเดียว แต่เป็น ecosystem ของ directional bets, arbitrage, macro, event, credit, quant และ multi-manager platforms", 38, TEXT, 1370, 1.18)
    for i, t in enumerate(["Thai deck", "16:9", "Data to Jun 2025 + FY2025 HFR", "Not investment advice"]):
        pill(d, 82 + i * 330, 545, t, [ACCENT, GREEN, AMBER, PURPLE][i], 290)
    draw_text(d, (82, 790), "Prepared from HFR, Aurum, and public manager strategy pages. ตัวเลขเป็นดัชนี/peer-group net returns ไม่ใช่ return ของกองทุนใดกองทุนหนึ่ง.", 26, MUTED, 1420)
    paths.append(save_slide(img, 1))

    img, d = base("Executive Summary", "Key findings")
    bullets = [
        ("1", "Multi-strategy ชนะเชิง risk-adjusted", "Aurum 5Y CAR ถึง Jun 2025: Multi-Strategy 11.1% ด้วย vol 2.7% และ Sharpe สูงสุดในกลุ่มหลัก"),
        ("2", "2022 คือ stress test ของโลก hedge fund", "Long-biased / Equity L/S / Event ติดลบ แต่ Macro, Quant และ Arbitrage ยังบวกหรือกัน drawdown ได้ดีกว่า"),
        ("3", "2025 กลับมาเป็นปีที่แรง", "HFR รายงาน HFRI Fund Weighted Composite +12.5% ในปี 2025 และ Equity Hedge +17.1%"),
        ("4", "กลยุทธ์ที่ดูเหมือนชื่อเดียวกันอาจ risk ต่างกันมาก", "ตัวอย่าง Event-driven มีทั้ง merger arb ต่ำกว่า beta, activist beta สูง, distressed credit, litigation และ special situations"),
    ]
    y = 250
    for n, title, body in bullets:
        d.ellipse((90, y + 8, 138, y + 56), fill=ACCENT)
        draw_text(d, (106, y + 16), n, 26, "#071018")
        draw_text(d, (160, y), title, 34, TEXT, 1200)
        y = draw_text(d, (160, y + 48), body, 25, MUTED, 1300, 1.18) + 34
    paths.append(save_slide(img, 2))

    img, d = base("Strategy Universe: กลยุทธ์ที่พบในอุตสาหกรรม", "Map")
    families = [
        ("Directional Equity", "Long/short equity, long-biased, short bias, sector specialist, equity market neutral, fundamental value/growth, healthcare/tech/energy, emerging market equity"),
        ("Event-Driven", "Merger arbitrage, activist, special situations, spin-off, index rebalance, capital structure catalyst, litigation, distressed/restructuring"),
        ("Relative Value / Arbitrage", "Convertible arb, fixed-income RV, volatility arb, statistical arb, pairs, basis, yield alternatives, asset-backed, mortgage RV, regulatory capital"),
        ("Macro / CTA / Systematic", "Global macro, discretionary macro, systematic diversified/CTA, trend following, commodities, rates, FX, emerging markets macro, risk premia, risk parity"),
        ("Credit / Specialty / Multi", "Credit L/S, distressed debt, direct lending, structured credit/CLO, insurance-linked securities, crypto, multi-strategy, pod shop, quant multi-asset"),
    ]
    for i, (title, body) in enumerate(families):
        card(d, (90 + (i % 2) * 870, 230 + (i // 2) * 245), (800, 190), title, body, [ACCENT, GREEN, AMBER, PURPLE, "#F97316"][i])
    paths.append(save_slide(img, 3))

    img, d = base("ใครใช้กลยุทธ์อะไรอยู่บ้าง", "Public examples")
    table_slide(d, strategy_examples, 90, 225, [520, 1160], 54, ["Strategy", "Publicly associated managers / platforms"], 20)
    draw_text(d, (92, 1015), "หมายเหตุ: เป็นตัวอย่างจาก public positioning/strategy pages และ industry coverage; ไม่ใช่รายชื่อครบทุกกองทุน และแต่ละ firm อาจใช้หลายกลยุทธ์พร้อมกัน.", 20, MUTED, 1580)
    paths.append(save_slide(img, 4))

    heat = IMG_DIR / "chart-major-returns.png"
    chart_major_heatmap(heat)
    img, d = base("5-Year Performance: กลุ่มหลัก", "Aurum net return to Jun 2025")
    d.rounded_rectangle((80, 210, 1840, 905), radius=18, fill="#071018")
    chart = Image.open(heat).resize((1640, 735))
    img.paste(chart, (145, 210))
    draw_text(d, (90, 940), "อ่านตาราง: 2025 เป็น H1/ถึง Jun 2025; 5Y CAR คือ compound annual return. Multi-Strategy เด่นสุดเมื่อเทียบ return กับ volatility.", 23, MUTED, 1560)
    paths.append(save_slide(img, 5))

    rr = IMG_DIR / "chart-risk-return.png"
    chart_risk_return(rr)
    img, d = base("Return vs Risk: ทำไม allocator ชอบ Multi-Strategy", "5Y CAR vs 5Y volatility")
    chart = Image.open(rr).resize((1160, 700))
    img.paste(chart, (100, 225))
    card(d, (1320, 250), (480, 180), "Core observation", "Multi-Strategy อยู่มุมซ้ายบน: return สูงและ vol ต่ำกว่า equity-heavy styles อย่างชัดเจน", GREEN)
    card(d, (1320, 470), (480, 180), "Tradeoff", "ผลลัพธ์มาจาก diversification, leverage, tight risk limits และ platform cost ที่สูงกว่า", AMBER)
    card(d, (1320, 690), (480, 180), "Due diligence", "ต้องดู crowding, liquidity, pass-through fees, netting risk และ key-person/platform risk", ACCENT)
    paths.append(save_slide(img, 6))

    sub = IMG_DIR / "chart-sub5.png"
    chart_sub5(sub)
    img, d = base("Sub-Strategy Winners: 5Y CAR สูงสุด", "Aurum to Jun 2025")
    chart = Image.open(sub).resize((1250, 780))
    img.paste(chart, (80, 210))
    card(d, (1390, 245), (420, 185), "Top cluster", "Activist, long-equity และ quant multi เป็นกลุ่มบนสุดในช่วง 5 ปี", GREEN)
    card(d, (1390, 475), (420, 185), "Cyclical context", "ผลลัพธ์ได้แรงหนุนจาก equity recovery, AI/tech leadership, commodity/rates dispersion", AMBER)
    card(d, (1390, 705), (420, 185), "Warning", "5Y winner ไม่ได้แปลว่าปีถัดไปชนะ; dispersion ระหว่างกองทุนสูงมาก", RED)
    paths.append(save_slide(img, 7))

    img, d = base("Directional Equity Family", "Long/short, market neutral, sector")
    rows = [
        ("Long/Short Equity", "ซื้อหุ้นที่มองว่าถูก/ดี และ short หุ้นที่แพง/แย่; net exposure มักเป็นบวก", "2022 เจ็บจาก equity beta; 2023-25 ฟื้นตามตลาดหุ้น"),
        ("Equity Market Neutral", "คุม beta ใกล้ศูนย์ ใช้ pairs, factors, statistical arb, sector neutral books", "return ต่ำกว่าแต่ drawdown มักนุ่มกว่า"),
        ("Sector Specialist", "โฟกัส tech, healthcare, energy, financials หรือ region", "ชนะมากเมื่อ sector dispersion สูง เช่น healthcare/energy/AI"),
        ("Short Bias", "net short หรือใช้ short เป็น alpha หลัก", "ดีใน crash แต่ถือยากใน bull market ยาว"),
    ]
    table_slide(d, rows, 90, 240, [390, 710, 560], 115, ["Sub-strategy", "How it works", "5Y read-through"], 22)
    paths.append(save_slide(img, 8))

    img, d = base("Event-Driven Family", "Catalyst investing")
    rows = [
        ("Merger Arb", "ซื้อ target / short acquirer หรือ hedge deal spread", "ดีเมื่อ M&A active และ financing/regulatory risk ต่ำ"),
        ("Activist", "ถือ stake แล้วผลักดัน buyback, spin-off, board change, margin improvement", "Aurum 5Y CAR 14.4% สูงสุดใน sub-strategy set"),
        ("Special Situations", "spin-off, index inclusion, litigation, holding company discount, recap", "alpha สูงแต่ idiosyncratic risk สูง"),
        ("Distressed / Restructuring", "ซื้อหนี้หรือหุ้นของบริษัทที่กำลัง restructure", "ดีใน credit cycle ที่มี forced sellers และ recovery value ชัด"),
    ]
    table_slide(d, rows, 90, 240, [390, 720, 550], 118, ["Sub-strategy", "How it works", "5Y read-through"], 22)
    paths.append(save_slide(img, 9))

    img, d = base("Relative Value / Arbitrage Family", "Mispricing, spreads, convergence")
    rows = [
        ("Convertible Arb", "long convertible bond + hedge equity/credit/rates Greeks", "Aurum 5Y CAR 7.8%; rate/vol regime สำคัญ"),
        ("Fixed Income RV", "เล่น curve, basis, swap spread, sovereign/cross-market relative value", "มี leverage สูง ต้องดู liquidity และ funding"),
        ("Volatility Arb", "เทียบ implied vs realized vol, skew, dispersion, options relative value", "ดีเมื่อ vol mispricing กว้าง แต่ short-vol tail risk ต้องชัด"),
        ("Stat Arb / Pairs", "โมเดล mean reversion และ factor-neutral equity baskets", "Aurum Quant-Stat Arb 5Y CAR 5.5% แต่ 1Y/3Y แข็งมาก"),
    ]
    table_slide(d, rows, 90, 240, [390, 720, 550], 118, ["Sub-strategy", "How it works", "5Y read-through"], 22)
    paths.append(save_slide(img, 10))

    img, d = base("Macro / CTA / Systematic Family", "Rates, FX, commodities, trends")
    rows = [
        ("Global Macro", "top-down view บน rates, FX, equity indices, commodities, policy/geopolitics", "2022 ดีเพราะ inflation/rates shock; 2023 softer; 2024-25 กลับมา"),
        ("CTA / Managed Futures", "systematic trend following across futures markets", "ชนะใน crisis ที่ trend ชัด เช่น 2022 inflation/rate rise"),
        ("Commodities Macro", "energy, metals, ags, power, weather/supply-demand", "Aurum 5Y CAR 10.4%; dispersion สูง"),
        ("Risk Premia / Risk Parity", "rules-based exposure to value, momentum, carry, defensive, balanced risk", "ขึ้นกับ regime correlation หุ้น/บอนด์"),
    ]
    table_slide(d, rows, 90, 240, [390, 720, 550], 118, ["Sub-strategy", "How it works", "5Y read-through"], 22)
    paths.append(save_slide(img, 11))

    img, d = base("Credit / Specialty / Digital Assets", "Less liquid, carry, convexity")
    rows = [
        ("Credit L/S", "long cheap credit / short weak credit; hedge duration or market beta", "Aurum Credit 5Y CAR 7.4%, vol 2.9%"),
        ("Distressed Credit", "ซื้อ debt ต่ำ par รอ restructuring/recovery", "Aurum 5Y CAR 8.5%; cycle-dependent"),
        ("Direct Lending", "private loans, spread/covenant/illiquidity premium", "Aurum 5Y CAR 6.9%; mark smoothing ต้องระวัง"),
        ("Crypto Hedge Funds", "long/short tokens, basis, market neutral, venture/liquid hybrid", "2021 boom, 2022 crash, 2023-25 rebound; volatility สูงมาก"),
    ]
    table_slide(d, rows, 90, 240, [390, 720, 550], 118, ["Sub-strategy", "How it works", "5Y read-through"], 22)
    paths.append(save_slide(img, 12))

    img, d = base("What Worked Over The Last 5 Years", "Interpretation")
    card(d, (90, 235), (520, 250), "Worked best", "Multi-strategy platforms, activist/event with strong catalysts, long-equity after 2022, commodity macro, selected quant multi.", GREEN)
    card(d, (700, 235), (520, 250), "Protected capital", "Macro/CTA, arbitrage และ credit ช่วยในปี 2022 เมื่อตลาดหุ้นและบอนด์พร้อมใจกันลง.", ACCENT)
    card(d, (1310, 235), (520, 250), "Underperformed risk-on beta", "กลยุทธ์ที่ hedge หนักจะตาม S&P 500 ไม่ทันในปี equity bull market เช่น 2023-2025.", AMBER)
    draw_text(d, (100, 575), "Practical allocation read:", 34, TEXT)
    practical = [
        "ใช้ Multi-strategy เป็น core diversifier ได้ แต่ต้องยอมรับ fee/pass-through และ capacity risk",
        "ใช้ Macro/CTA เป็น crisis convexity และ inflation/rates shock hedge",
        "ใช้ Equity L/S/Event เป็น alpha sleeve ที่ยังพึ่ง equity cycle บางส่วน",
        "ใช้ Credit/Direct Lending เมื่อรับ illiquidity ได้ และต้อง stress test default/recovery",
    ]
    y = 635
    for b in practical:
        d.ellipse((110, y + 10, 126, y + 26), fill=ACCENT)
        y = draw_text(d, (150, y), b, 30, MUTED, 1500, 1.16) + 24
    paths.append(save_slide(img, 13))

    img, d = base("Sources & Methodology", "Research notes")
    notes = [
        "Performance table: Aurum Hedge Fund Data Engine, net returns, 5-year period to June 2025.",
        "Full-year 2025 headline: HFR reported HFRI Fund Weighted Composite +12.5%, Equity Hedge +17.1%, Event-Driven +10.9%.",
        "2024/2023/2022/2021 HFR notes used to cross-check main strategy behavior and market narratives.",
        "Manager examples come from public strategy pages and industry-known public positioning; individual fund holdings/returns are not inferred.",
        "This is research material, not investment advice. Hedge fund indices have reporting, survivorship, liquidity and selection biases.",
    ]
    y = 230
    for n in notes:
        d.rectangle((90, y + 9, 104, y + 23), fill=ACCENT)
        y = draw_text(d, (130, y), n, 27, TEXT, 1500, 1.18) + 28
    y += 20
    draw_text(d, (90, y), "Primary sources:", 30, ACCENT)
    y += 46
    for name, url in sources[:9]:
        y = draw_text(d, (110, y), f"{name} — {url}", 19, MUTED, 1640, 1.08) + 8
    paths.append(save_slide(img, 14))

    return paths


def make_pptx(slide_paths):
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for path in slide_paths:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width, height=prs.slide_height)
    out = OUT_DIR / f"{PROJECT}.pptx"
    prs.save(out)
    return out


def write_research():
    md = NOTES_DIR / f"{PROJECT}.md"
    lines = [
        "# Hedge Fund Strategies Research",
        "",
        "ภาษา: ไทย | Scope: กลยุทธ์ hedge fund, ตัวอย่างผู้ใช้, และ performance 5 ปีย้อนหลัง",
        "",
        "## Executive summary",
        "",
        "- Hedge fund เป็น universe ของกลยุทธ์ ไม่ใช่ asset class เดียว กลุ่มหลักคือ Equity, Event-Driven, Relative Value/Arbitrage, Macro/CTA, Credit/Specialty และ Multi-Strategy",
        "- ช่วง 5 ปีถึง Jun 2025 ตาม Aurum: Multi-Strategy เด่นสุดในกลุ่มหลักด้วย 5Y CAR 11.08% และ volatility 2.69%",
        "- ปี 2022 เป็นปีคัดแยกกลยุทธ์: equity/event/long-biased ติดลบ แต่ macro, quant, arbitrage และบาง credit sleeves ช่วยกัน downside ได้",
        "- ปี 2025 เต็มปีตาม HFR แข็งแรงมาก: HFRI Fund Weighted Composite +12.5%, Equity Hedge +17.1%, Event-Driven +10.9%",
        "",
        "## Major strategy performance",
        "",
        major_returns.to_markdown(index=False),
        "",
        "หมายเหตุ: Aurum 2025 เป็น H1/ถึง June 2025; 5Y CAR/Vol เป็นช่วง 5 ปีถึง June 2025",
        "",
        "## Sub-strategy 5Y CAR ranking",
        "",
        sub_5y.to_markdown(index=False),
        "",
        "## Strategy taxonomy and example users",
        "",
    ]
    for strat, users in strategy_examples:
        lines.append(f"- **{strat}:** {users}")
    lines += [
        "",
        "## Full strategy list found",
        "",
        "Long/short equity; long-biased equity; short bias; equity market neutral; statistical arbitrage; quantitative directional; sector specialist; fundamental value; fundamental growth; emerging-market equity; merger arbitrage; activist; special situations; spin-offs; index rebalance; litigation finance/event; distressed/restructuring; capital structure arbitrage; convertible arbitrage; fixed income relative value; mortgage/asset-backed relative value; volatility arbitrage; options dispersion; yield alternatives; credit long/short; distressed credit; direct lending; structured credit/CLO; municipal credit; global macro; discretionary macro; systematic macro; CTA/managed futures; trend following; commodity macro; FX/rates macro; emerging-market macro; risk premia; risk parity; multi-strategy; multi-manager pod shop; crypto/digital asset long-short; crypto basis/market-neutral; insurance-linked securities/reinsurance; regulatory capital trades.",
        "",
        "## Source list",
        "",
    ]
    for name, url in sources:
        lines.append(f"- {name}: {url}")
    md.write_text("\n".join(lines), encoding="utf-8")
    return md


if __name__ == "__main__":
    slide_paths = make_slides()
    pptx = make_pptx(slide_paths)
    md = write_research()
    print(pptx)
    print(md)
    print(IMG_DIR)
