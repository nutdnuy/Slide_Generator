from pathlib import Path
import textwrap

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
PROJECT = "chat-thinks-agent-does"
IMG_DIR = ROOT / "assets" / "images" / PROJECT
OUT_DIR = ROOT / "outputs"
NOTES_DIR = ROOT / "notes"

W, H = 1920, 1080

BG = "#0B0F14"
PANEL = "#131A22"
PANEL_2 = "#172330"
LINE = "#2B3948"
TEXT = "#F6F4EF"
MUTED = "#AEB7C2"
ORANGE = "#D97742"
TEAL = "#43C6AC"
BLUE = "#6AA7FF"
PURPLE = "#9D7CFF"
GREEN = "#72D67C"
AMBER = "#F5C356"
RED = "#EA6A6A"

FONT_REGULAR = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"


def ensure_dirs():
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    NOTES_DIR.mkdir(parents=True, exist_ok=True)


def font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, size)


def text_bbox(draw, text, f):
    return draw.textbbox((0, 0), text, font=f)


def wrap_text(draw, text, f, max_width):
    lines = []
    for paragraph in str(text).split("\n"):
        words = paragraph.split(" ")
        line = ""
        for word in words:
            candidate = (line + " " + word).strip()
            if text_bbox(draw, candidate, f)[2] <= max_width or not line:
                line = candidate
            else:
                lines.append(line)
                line = word
        lines.append(line)
    return lines


def draw_text(draw, xy, text, size=34, fill=TEXT, max_width=None, bold=False, spacing=1.16):
    f = font(size, bold)
    x, y = xy
    if max_width is None:
        draw.text((x, y), text, font=f, fill=fill)
        return y + size
    for line in wrap_text(draw, text, f, max_width):
        draw.text((x, y), line, font=f, fill=fill)
        y += int(size * spacing)
    return y


def pill(draw, x, y, text, fill, fg="#071018", pad_x=24, pad_y=10, size=24):
    f = font(size, True)
    box = text_bbox(draw, text, f)
    w = box[2] - box[0] + pad_x * 2
    h = box[3] - box[1] + pad_y * 2 + 2
    draw.rounded_rectangle((x, y, x + w, y + h), radius=h // 2, fill=fill)
    draw.text((x + pad_x, y + pad_y - 1), text, font=f, fill=fg)
    return x + w


def rounded_panel(draw, xy, radius=28, fill=PANEL, outline=LINE, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def draw_header(draw, title, subtitle=None):
    draw.rectangle((0, 0, W, H), fill=BG)
    draw.rectangle((0, 0, W, 86), fill="#070A0E")
    draw.line((80, 86, W - 80, 86), fill="#24303C", width=2)
    pill(draw, 80, 25, "CLAUDE PRODUCT SHAPES", ORANGE, size=19)
    draw_text(draw, (80, 132), title, 66, TEXT, W - 160, True, 1.05)
    if subtitle:
        draw_text(draw, (82, 218), subtitle, 31, MUTED, W - 164, False, 1.18)


def draw_person_chat(draw, x, y):
    draw.ellipse((x + 82, y + 20, x + 178, y + 116), fill="#F0C9A2")
    draw.arc((x + 64, y + 8, x + 196, y + 126), 190, 350, fill="#6D4B35", width=20)
    draw.rounded_rectangle((x + 54, y + 122, x + 206, y + 316), radius=42, fill=BLUE)
    draw.rectangle((x + 86, y + 288, x + 174, y + 430), fill="#26384D")
    draw.line((x + 54, y + 190, x - 28, y + 282), fill="#F0C9A2", width=22)
    draw.line((x + 206, y + 190, x + 302, y + 266), fill="#F0C9A2", width=22)
    draw.rounded_rectangle((x + 280, y + 92, x + 628, y + 248), radius=28, fill="#F8F3EA")
    draw.polygon([(x + 312, y + 236), (x + 260, y + 284), (x + 340, y + 252)], fill="#F8F3EA")
    draw_text(draw, (x + 316, y + 125), "Claude,\nhelp me think\nthrough this.", 32, "#18202A", 260, True, 1.08)
    draw.rounded_rectangle((x + 260, y + 318, x + 650, y + 420), radius=20, fill="#1F2A36", outline="#34485F", width=2)
    draw_text(draw, (x + 292, y + 348), "Turn 1  ->  Turn 2  ->  Turn 3", 31, MUTED, 326, False)


def draw_agent_tools(draw, x, y):
    # Claude agent core
    draw.rounded_rectangle((x + 246, y + 124, x + 478, y + 356), radius=44, fill="#EFE7DA", outline=ORANGE, width=5)
    draw.ellipse((x + 312, y + 184, x + 412, y + 284), outline="#161A1E", width=8)
    draw.line((x + 362, y + 150, x + 362, y + 330), fill="#161A1E", width=6)
    draw.line((x + 278, y + 234, x + 446, y + 234), fill="#161A1E", width=6)
    draw_text(draw, (x + 305, y + 374), "Agent loop", 26, TEXT, 160, True)

    tools = [
        ("bash", x + 20, y + 18, TEAL),
        ("browser", x + 510, y + 32, BLUE),
        ("files", x + 24, y + 394, AMBER),
        ("MCP", x + 524, y + 392, PURPLE),
    ]
    for label, tx, ty, color in tools:
        draw.line((x + 362, y + 238, tx + 122, ty + 62), fill=color, width=5)
        draw.rounded_rectangle((tx, ty, tx + 244, ty + 124), radius=22, fill="#101820", outline=color, width=3)
        if label == "bash":
            draw_text(draw, (tx + 26, ty + 28), "$ run task", 29, color, 190, True)
            draw_text(draw, (tx + 26, ty + 68), "tests passed", 22, MUTED, 180)
        elif label == "browser":
            draw.rectangle((tx + 24, ty + 28, tx + 220, ty + 50), fill="#223043")
            draw_text(draw, (tx + 26, ty + 66), "inspect UI", 27, color, 184, True)
        elif label == "files":
            draw.polygon([(tx + 32, ty + 26), (tx + 142, ty + 26), (tx + 196, ty + 78), (tx + 196, ty + 104), (tx + 32, ty + 104)], fill="#263241")
            draw_text(draw, (tx + 58, ty + 60), "edit files", 25, color, 150, True)
        else:
            draw_text(draw, (tx + 34, ty + 32), "connect", 29, color, 170, True)
            draw_text(draw, (tx + 34, ty + 70), "external tools", 21, MUTED, 170)


def slide_1():
    img = Image.new("RGB", (W, H), BG)
    draw = ImageDraw.Draw(img)
    draw_header(draw, "Chat Thinks · Agent Does", "Claude ships in two product shapes that solve different problems")

    left = (80, 310, 910, 900)
    right = (1010, 310, 1840, 900)
    rounded_panel(draw, left, fill="#101720", outline="#26384A")
    rounded_panel(draw, right, fill="#111A20", outline="#2A4039")
    pill(draw, 130, 344, "Driving manually", BLUE, size=25)
    pill(draw, 1060, 344, "Autopilot", TEAL, size=25)
    draw_text(draw, (130, 410), "Chat = user steers each turn", 37, TEXT, 680, True)
    draw_text(draw, (1060, 410), "Agent = goal runs through tools", 37, TEXT, 680, True)
    draw_person_chat(draw, 150, 452)
    draw_agent_tools(draw, 1090, 450)
    draw.line((960, 326, 960, 886), fill="#34475A", width=3)
    draw_text(draw, (878, 558), "vs", 56, "#6D7783", 120, True)
    return img


def slide_2():
    img = Image.new("RGB", (W, H), BG)
    draw = ImageDraw.Draw(img)
    draw_header(draw, "Chat vs Agent — 6 Dimensions That Matter")

    x0, y0 = 92, 282
    widths = [420, 635, 635]
    row_h = 104
    headers = ["Dimension", "Chat", "Agent"]
    fills = ["#111820", "#142033", "#13251F"]

    draw.rounded_rectangle((x0, y0, x0 + sum(widths), y0 + row_h * 7), radius=28, fill="#0F151C", outline=LINE, width=2)
    cx = x0
    for i, (hdr, w) in enumerate(zip(headers, widths)):
        draw.rectangle((cx, y0, cx + w, y0 + row_h), fill=fills[i])
        draw_text(draw, (cx + 28, y0 + 32), hdr, 34, TEXT, w - 56, True)
        cx += w
    draw.line((x0, y0 + row_h, x0 + sum(widths), y0 + row_h), fill="#34465A", width=3)

    rows = [
        ("Interaction", "Turn-based, synchronous", "Goal-based, long-running"),
        ("Autonomy", "User-in-the-loop", "Human-on-the-loop"),
        ("Tools", "Limited: web, files, artifacts", "Open-ended: MCP, code, browser, OS"),
        ("Time horizon", "Seconds–minutes", "Minutes–hours"),
        ("Output", "An answer", "A finished task"),
        ("Product surface", "Claude.ai", "Claude Code, Agent SDK, Computer Use"),
    ]

    y = y0 + row_h
    for r, row in enumerate(rows):
        if r % 2 == 1:
            draw.rectangle((x0, y, x0 + sum(widths), y + row_h), fill="#101820")
        cx = x0
        for c, (txt, w) in enumerate(zip(row, widths)):
            fill = TEXT if c == 0 else ("#D6E6FF" if c == 1 else "#D8F5EA")
            draw_text(draw, (cx + 28, y + 30), txt, 27 if c else 29, fill, w - 56, c == 0, 1.08)
            if c < 2:
                draw.line((cx + w, y, cx + w, y + row_h), fill="#263645", width=2)
            cx += w
        draw.line((x0, y + row_h, x0 + sum(widths), y + row_h), fill="#202D39", width=2)
        y += row_h

    pill(draw, 1060, 218, "Agent shifts Claude from response surface to execution surface", ORANGE, size=22)
    return img


def bullet_list(draw, x, y, items, color, max_width):
    f = font(28, False)
    for item in items:
        draw.ellipse((x, y + 9, x + 12, y + 21), fill=color)
        lines = wrap_text(draw, item, f, max_width - 34)
        for j, line in enumerate(lines):
            draw.text((x + 34, y), line, font=f, fill=TEXT)
            y += 34 if j < len(lines) - 1 else 44
    return y


def value_block(draw, x, y, label, value, color):
    draw_text(draw, (x, y), label, 23, color, 260, True)
    return draw_text(draw, (x, y + 30), value, 31, TEXT, 650, True, 1.08) + 18


def slide_3():
    img = Image.new("RGB", (W, H), BG)
    draw = ImageDraw.Draw(img)
    draw_header(draw, "Where Each Wins")

    left = (96, 274, 908, 850)
    right = (1012, 274, 1824, 850)
    rounded_panel(draw, left, fill="#111A28", outline="#294568")
    rounded_panel(draw, right, fill="#101E19", outline="#2B5548")

    draw_text(draw, (142, 326), "Chat → Knowledge Work", 45, "#DDEBFF", 700, True)
    draw_text(draw, (1058, 326), "Agent → Execution Work", 45, "#DDF7EA", 700, True)

    y_left = bullet_list(
        draw,
        146,
        416,
        ["Q&A", "Summarization", "Drafting", "Brainstorming", "Explanation"],
        BLUE,
        660,
    )
    y_right = bullet_list(
        draw,
        1062,
        416,
        ["Code migration", "KYC review", "Research pipelines", "End-to-end automation"],
        TEAL,
        660,
    )

    y_left = max(y_left + 10, 654)
    y_right = max(y_right + 10, 654)
    value_block(draw, 146, y_left, "VALUE", "Reduces time-to-answer", AMBER)
    value_block(draw, 1062, y_right, "VALUE", "Reduces headcount cost on repetitive workflows", AMBER)
    value_block(draw, 146, 752, "SOLD AS", "Productivity Tool (per-seat)", ORANGE)
    value_block(draw, 1062, 752, "SOLD AS", "Digital Workforce (per-task)", ORANGE)

    draw.rounded_rectangle((170, 902, 1750, 994), radius=28, fill="#EDE5D6")
    draw_text(draw, (225, 928), "Chat is Claude that answers. Agent is Claude that ships.", 47, "#12171D", 1470, True)
    return img


def save_slides():
    slides = [slide_1(), slide_2(), slide_3()]
    paths = []
    for i, img in enumerate(slides, 1):
        path = IMG_DIR / f"slide-{i:02d}.png"
        img.save(path, quality=96)
        paths.append(path)
    return paths


def build_pptx(paths):
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for path in paths:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width, height=prs.slide_height)
    # Remove default empty slide if present in template.
    if len(prs.slides) > len(paths):
        xml_slides = prs.slides._sldIdLst
        rel_id = xml_slides[0].rId
        prs.part.drop_rel(rel_id)
        xml_slides.remove(xml_slides[0])
    out = OUT_DIR / f"{PROJECT}.pptx"
    prs.save(out)
    return out


def contact_sheet(paths):
    thumbs = []
    for p in paths:
        im = Image.open(p).resize((640, 360))
        thumbs.append(im)
    sheet = Image.new("RGB", (640 * 3 + 40, 400), "#05070A")
    d = ImageDraw.Draw(sheet)
    x = 10
    for idx, im in enumerate(thumbs, 1):
        sheet.paste(im, (x, 10))
        draw_text(d, (x + 12, 370), f"Slide {idx}", 20, MUTED, 120, True)
        x += 650
    path = NOTES_DIR / f"{PROJECT}-contact-sheet.png"
    sheet.save(path, quality=95)
    return path


def write_outline():
    outline = """# Chat Thinks · Agent Does

Audience: Business/product stakeholders comparing Claude chat and agent product surfaces.
Objective: Explain when to use chat versus agent, and why the business value model changes.
Language: English
Slide count: 3
Aspect ratio: 16:9
Rendering mode: Deterministic full-slide PNGs embedded in PPTX for exact text layout.

## Slide 1 — The Big Idea (Hook)
Message: Chat helps a user think turn by turn; Agent executes a goal through tools.

## Slide 2 — Side-by-Side Comparison
Message: The difference is not just UI. Autonomy, tool access, and time horizon change the product category.

## Slide 3 — When to Use & Business Value
Message: Chat wins knowledge work; Agent wins execution work and can be sold as digital workforce capacity.
"""
    path = NOTES_DIR / f"{PROJECT}-outline.md"
    path.write_text(outline, encoding="utf-8")
    return path


def main():
    ensure_dirs()
    paths = save_slides()
    pptx = build_pptx(paths)
    sheet = contact_sheet(paths)
    outline = write_outline()
    print(f"PPTX: {pptx}")
    print(f"Slides: {IMG_DIR}")
    print(f"Contact sheet: {sheet}")
    print(f"Outline: {outline}")


if __name__ == "__main__":
    main()
