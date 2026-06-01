from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches


def main():
    if len(sys.argv) != 3:
        print("Usage: python scripts/build_image_deck.py <project> <count>", file=sys.stderr)
        sys.exit(2)

    root = Path(__file__).resolve().parents[1]
    project = sys.argv[1]
    count = int(sys.argv[2])
    image_dir = root / "assets" / "images" / project / "final"
    out_dir = root / "outputs"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{project}.pptx"

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for i in range(1, count + 1):
        image = image_dir / f"slide-{i:02d}.png"
        if not image.exists():
            raise FileNotFoundError(image)
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(image), 0, 0, width=prs.slide_width, height=prs.slide_height)

    if len(prs.slides) > count:
        xml_slides = prs.slides._sldIdLst
        rel_id = xml_slides[0].rId
        prs.part.drop_rel(rel_id)
        xml_slides.remove(xml_slides[0])

    prs.save(out_path)
    print(out_path)


if __name__ == "__main__":
    main()
