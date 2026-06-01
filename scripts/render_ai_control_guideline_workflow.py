from __future__ import annotations

import json
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageEnhance, ImageFont
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
PROJECT = "ai-control-guideline-workflow"
IMAGE_ROOT = ROOT / "assets" / "images" / PROJECT
ORIGINALS = IMAGE_ROOT / "originals"
FINAL = IMAGE_ROOT / "final"
OUT_PPTX = ROOT / "outputs" / f"{PROJECT}.pptx"
OUTLINE = ROOT / "notes" / f"{PROJECT}-outline.md"
PLAN_JSON = ROOT / "notes" / f"{PROJECT}-plan.json"
CONTACT_SHEET = IMAGE_ROOT / "contact-sheet.png"

W, H = 1920, 1080
SLIDE_W, SLIDE_H = 13.333333, 7.5

# Tahoma renders Thai vowels and tone marks correctly through Pillow's basic
# text engine in this runtime. Thonburi renders some combining marks as tofu.
FONT_REG = "/System/Library/Fonts/Supplemental/Tahoma.ttf"
FONT_FALLBACK = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"

C = {
    "ink": "#514846",
    "ink2": "#2F2B2A",
    "muted": "#786F6C",
    "line": "#DDD6D1",
    "paper": "#FFFFFF",
    "soft": "#F6F3EF",
    "soft_yellow": "#FFF5CF",
    "yellow": "#FFC800",
    "yellow2": "#FFD84A",
    "taupe": "#3B3432",
    "green": "#2E7D32",
    "blue": "#315E7D",
    "red": "#B3261E",
}


def load_font(size: int) -> ImageFont.FreeTypeFont:
    try:
        return ImageFont.truetype(FONT_REG, size=size)
    except OSError:
        return ImageFont.truetype(FONT_FALLBACK, size=size)


FONTS = {size: load_font(size) for size in range(16, 76)}


def f(size: int) -> ImageFont.FreeTypeFont:
    size = max(16, min(75, int(size)))
    return FONTS[size]


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", str(text).replace("\u200b", " ")).strip()


def measure(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont) -> int:
    if not text:
        return 0
    bbox = draw.textbbox((0, 0), text, font=font)
    return bbox[2] - bbox[0]


def split_long_token(draw: ImageDraw.ImageDraw, token: str, font: ImageFont.FreeTypeFont, max_w: int) -> list[str]:
    parts: list[str] = []
    current = ""
    for ch in token:
        candidate = current + ch
        if current and measure(draw, candidate, font) > max_w:
            parts.append(current)
            current = ch
        else:
            current = candidate
    if current:
        parts.append(current)
    return parts


def wrap_lines(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, max_w: int) -> list[str]:
    lines: list[str] = []
    for paragraph in str(text).split("\n"):
        paragraph = paragraph.strip()
        if not paragraph:
            lines.append("")
            continue
        words = paragraph.split(" ")
        current = ""
        for word in words:
            candidate = word if not current else f"{current} {word}"
            if measure(draw, candidate, font) <= max_w:
                current = candidate
                continue
            if current:
                lines.append(current)
                current = ""
            if measure(draw, word, font) <= max_w:
                current = word
            else:
                chunks = split_long_token(draw, word, font, max_w)
                lines.extend(chunks[:-1])
                current = chunks[-1] if chunks else ""
        if current:
            lines.append(current)
    return lines


def draw_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    x: int,
    y: int,
    max_w: int,
    font: ImageFont.FreeTypeFont,
    fill: str,
    *,
    line_gap: int = 8,
    max_lines: int | None = None,
    bold: bool = False,
    center: bool = False,
) -> int:
    lines = wrap_lines(draw, text, font, max_w)
    if max_lines is not None and len(lines) > max_lines:
        lines = lines[:max_lines]
        if lines:
            while lines[-1] and measure(draw, lines[-1] + "...", font) > max_w:
                lines[-1] = lines[-1][:-1]
            lines[-1] = lines[-1].rstrip() + "..."
    cursor = y
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        line_w = bbox[2] - bbox[0]
        draw_x = x + (max_w - line_w) // 2 if center else x
        offsets = [(0, 0), (1, 0)] if bold else [(0, 0)]
        for dx, dy in offsets:
            draw.text((draw_x + dx, cursor + dy), line, font=font, fill=fill)
        cursor += bbox[3] - bbox[1] + line_gap
    return cursor


def rounded(draw: ImageDraw.ImageDraw, box, fill, outline=None, width=1, radius=22):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def soften_background(img: Image.Image, amount: float = 1.05) -> Image.Image:
    img = img.resize((W, H), Image.Resampling.LANCZOS).convert("RGBA")
    return ImageEnhance.Brightness(img).enhance(amount)


def footer(draw: ImageDraw.ImageDraw, slide_no: int):
    draw_text(
        draw,
        "Source: Control_Guideline sheet | AI Risk Assessment v2 and Minimum Control Checklist",
        84,
        1015,
        1220,
        f(20),
        C["muted"],
        max_lines=1,
    )
    draw_text(draw, f"{slide_no:02d} / 13", 1670, 1015, 170, f(22), C["muted"], center=True)


def header(draw: ImageDraw.ImageDraw, title: str, subtitle: str, slide_no: int, tag: str):
    draw.rectangle((0, 0, 18, H), fill=C["yellow"])
    rounded(draw, (1510, 48, 1842, 92), C["taupe"], radius=22)
    draw_text(draw, tag, 1530, 57, 292, f(20), "#FFFFFF", bold=True, center=True)
    draw_text(draw, "AI-COE | Control Guideline", 1490, 104, 360, f(18), C["muted"], center=True)
    draw_text(draw, title, 78, 50, 1330, f(48), C["ink2"], bold=True, max_lines=2, line_gap=4)
    draw_text(draw, subtitle, 82, 138, 1350, f(24), C["muted"], max_lines=2, line_gap=5)
    footer(draw, slide_no)


def badge(draw: ImageDraw.ImageDraw, text: str, x: int, y: int, fill: str, fg: str = "#FFFFFF", w: int | None = None):
    font = f(21)
    pad_x = 22
    width = w or measure(draw, text, font) + pad_x * 2
    rounded(draw, (x, y, x + width, y + 42), fill, radius=21)
    draw_text(draw, text, x + pad_x, y + 8, width - pad_x * 2, font, fg, bold=True, center=True, max_lines=1)
    return width


def draw_chips(draw: ImageDraw.ImageDraw, labels: list[str], x: int, y: int, max_w: int, *, fill=C["soft"], fg=C["ink"]):
    cx, cy = x, y
    font = f(22)
    for label in labels:
        label = clean(label)
        chip_w = min(max_w, measure(draw, label, font) + 34)
        if cx + chip_w > x + max_w:
            cx = x
            cy += 50
        rounded(draw, (cx, cy, cx + chip_w, cy + 38), fill, outline=C["line"], radius=18)
        draw_text(draw, label, cx + 17, cy + 7, chip_w - 34, font, fg, max_lines=1)
        cx += chip_w + 12


def draw_steps_panel(
    draw: ImageDraw.ImageDraw,
    steps: list[tuple[str, str]],
    *,
    x: int = 72,
    y: int = 278,
    w: int = 1040,
    h: int = 650,
):
    rounded(draw, (x, y, x + w, y + h), (255, 255, 255, 238), outline=C["line"], width=2, radius=28)
    draw.rectangle((x, y, x + 12, y + h), fill=C["yellow"])
    draw_text(draw, "ทีมโปรเจกต์ต้องทำอะไร", x + 34, y + 24, w - 72, f(34), C["ink2"], bold=True, max_lines=1)
    top = y + 84
    gap = 14
    card_h = int((h - 112 - gap * (len(steps) - 1)) / len(steps))
    for i, (head, detail) in enumerate(steps, start=1):
        cy = top + (i - 1) * (card_h + gap)
        rounded(draw, (x + 30, cy, x + w - 34, cy + card_h), C["soft"], outline="#ECE7E2", radius=20)
        circle = (x + 52, cy + 22, x + 102, cy + 72)
        draw.ellipse(circle, fill=C["yellow"])
        draw_text(draw, str(i), x + 52, cy + 31, 50, f(24), C["ink2"], bold=True, center=True)
        text_x = x + 124
        draw_text(draw, head, text_x, cy + 16, w - 190, f(26), C["ink2"], bold=True, max_lines=1)
        draw_text(draw, detail, text_x, cy + 52, w - 190, f(22), C["muted"], max_lines=2, line_gap=3)


def draw_right_panels(
    draw: ImageDraw.ImageDraw,
    narrative: str,
    evidence: list[str],
    phases: list[str],
    *,
    x: int = 1150,
    y: int = 278,
    w: int = 700,
):
    rounded(draw, (x, y, x + w, y + 270), C["taupe"], radius=28)
    draw_text(draw, "เล่าให้ทีมฟังแบบง่าย", x + 34, y + 28, w - 68, f(30), C["yellow"], bold=True, max_lines=1)
    draw_text(draw, narrative, x + 36, y + 82, w - 72, f(26), "#FFFFFF", max_lines=5, line_gap=6)

    rounded(draw, (x, y + 308, x + w, y + 650), (255, 245, 207, 246), outline=C["line"], width=2, radius=28)
    draw_text(draw, "Evidence ที่ต้องเตรียม", x + 34, y + 332, w - 68, f(30), C["ink2"], bold=True, max_lines=1)
    draw_chips(draw, evidence, x + 34, y + 388, w - 68, fill="#FFFFFF", fg=C["ink2"])

    if phases:
        draw_text(draw, "AI-SLC phase / gate", x + 34, y + 560, w - 68, f(22), C["muted"], bold=True, max_lines=1)
        draw_chips(draw, phases, x + 34, y + 600, w - 68, fill="#F7EFE3", fg=C["ink2"])


def render_cover(bg: Image.Image) -> Image.Image:
    img = soften_background(bg, 1.02)
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, 18, H), fill=C["yellow"])
    # Dark generated panel on slide 01 is kept and used for the title.
    draw_text(draw, "AI Control\nGuideline Workflow", 110, 205, 820, f(64), "#FFFFFF", bold=True, line_gap=8)
    draw_text(
        draw,
        "Applied condition for project > ต้องทำอะไร > ต้องเก็บ evidence อะไร",
        116,
        405,
        780,
        f(31),
        "#F8F2EA",
        max_lines=3,
        line_gap=8,
    )
    badge(draw, "Applicable = TRUE", 116, 545, C["yellow"], C["ink2"])
    badge(draw, "21 controls", 390, 545, "#FFFFFF", C["ink2"])
    badge(draw, "13 image slides", 610, 545, "#FFFFFF", C["ink2"])

    rounded(draw, (1030, 160, 1815, 840), (255, 255, 255, 232), outline=C["line"], radius=30)
    draw_text(draw, "ใช้เล่าให้คนที่ต้องทำเข้าใจเร็ว", 1080, 205, 700, f(38), C["ink2"], bold=True, max_lines=2)
    draw_text(draw, "Deck นี้สรุปเฉพาะ control ที่ Applicable เป็น TRUE ใน sheet Control_Guideline แล้วแปลงเป็น workflow ที่ทีมโปรเจกต์ทำตามได้", 1082, 285, 670, f(27), C["muted"], max_lines=5, line_gap=6)
    draw_chips(
        draw,
        ["Initiation", "Development", "Implementation", "Monitoring", "Evidence handoff"],
        1082,
        535,
        660,
        fill=C["soft_yellow"],
        fg=C["ink2"],
    )
    draw_text(draw, "Style reference: AISLC.pptx", 1084, 780, 650, f(23), C["muted"], max_lines=1)
    footer(draw, 1)
    return img


def render_overview(bg: Image.Image) -> Image.Image:
    img = soften_background(bg)
    draw = ImageDraw.Draw(img)
    header(draw, "วิธีใช้ deck นี้เป็น workflow", "เริ่มจาก condition แล้วไล่ไปที่ action, evidence และ phase/gate ที่เกี่ยวข้อง", 2, "HOW TO USE")
    labels = [
        ("1", "เช็คเงื่อนไข", "โปรเจกต์เข้า condition ไหนบ้าง"),
        ("2", "เปิด control", "ดูว่ามี action อะไรที่ต้องทำ"),
        ("3", "ทำตาม phase", "เตรียมงานตาม AI-SLC phase"),
        ("4", "ส่ง evidence", "รวบรวม artifact ให้ reviewer ตรวจได้"),
    ]
    x0, y0, gap = 118, 272, 28
    card_w, card_h = 400, 390
    for i, (num, head, detail) in enumerate(labels):
        x = x0 + i * (card_w + gap)
        rounded(draw, (x, y0, x + card_w, y0 + card_h), (255, 255, 255, 238), outline=C["line"], width=2, radius=30)
        draw.ellipse((x + 38, y0 + 42, x + 110, y0 + 114), fill=C["yellow"])
        draw_text(draw, num, x + 38, y0 + 56, 72, f(34), C["ink2"], bold=True, center=True)
        draw_text(draw, head, x + 42, y0 + 150, card_w - 84, f(35), C["ink2"], bold=True, center=True, max_lines=2)
        draw_text(draw, detail, x + 50, y0 + 238, card_w - 100, f(27), C["muted"], center=True, max_lines=3, line_gap=5)
        if i < 3:
            draw.line((x + card_w + 8, y0 + 195, x + card_w + gap - 8, y0 + 195), fill=C["yellow"], width=6)
            draw.polygon([(x + card_w + gap - 8, y0 + 195), (x + card_w + gap - 28, y0 + 182), (x + card_w + gap - 28, y0 + 208)], fill=C["yellow"])
    rounded(draw, (350, 770, 1570, 900), C["taupe"], radius=28)
    draw_text(draw, "Rule of thumb: ถ้า condition เป็น Yes / Applicable ให้ทำ action ทั้งหมดในกลุ่มนั้น และเก็บ evidence ก่อนเข้า gate", 410, 807, 1100, f(31), "#FFFFFF", center=True, max_lines=2)
    return img


def render_condition_map(bg: Image.Image) -> Image.Image:
    img = soften_background(bg)
    draw = ImageDraw.Draw(img)
    header(draw, "Applicable conditions ทั้งหมดที่ต้องเล่า", "สรุปจากแถวที่คอลัมน์ Applicable = TRUE ใน sheet Control_Guideline", 3, "CONDITION MAP")
    cards = [
        ("ทุกกรณี", "11 controls", "Baseline controls ที่ทุก AI use case ต้องทำ"),
        ("AI Risk = Medium or High", "1 control", "ต้องทดสอบ edge case"),
        ("AI Risk = High", "5 controls", "ต้องมี robustness, alert, kill switch และแผนหยุดใช้"),
        ("งานหลัก / กระทบลูกค้า", "1 control", "ต้องออกแบบ HITL / HOTL oversight"),
        ("Business-critical supervised model", "1 control", "ต้องทำ explainability ด้วย SHAP / LIME / feature importance"),
        ("AI สนทนากับลูกค้า", "1 control", "ต้องมี disclosure และ screenshot evidence"),
        ("External open data", "1 control", "ต้องบันทึก provenance และ license ใน ARB Pack"),
    ]
    x0, y0 = 88, 250
    col_w, row_h = 840, 135
    for idx, (title, count, desc) in enumerate(cards):
        col = idx % 2
        row = idx // 2
        x = x0 + col * (col_w + 62)
        y = y0 + row * (row_h + 24)
        rounded(draw, (x, y, x + col_w, y + row_h), (255, 255, 255, 240), outline=C["line"], width=2, radius=24)
        draw.rectangle((x, y, x + 12, y + row_h), fill=C["yellow"])
        draw_text(draw, title, x + 34, y + 22, 560, f(29), C["ink2"], bold=True, max_lines=1)
        rounded(draw, (x + 620, y + 22, x + 800, y + 61), C["soft_yellow"], outline=C["line"], radius=19)
        draw_text(draw, count, x + 635, y + 29, 150, f(21), C["ink2"], bold=True, center=True, max_lines=1)
        draw_text(draw, desc, x + 34, y + 72, 735, f(22), C["muted"], max_lines=2)
    rounded(draw, (530, 890, 1390, 958), C["taupe"], radius=28)
    draw_text(draw, "รวม 21 applicable controls > แปลงเป็น 10 workflow slides ถัดไป", 570, 906, 780, f(28), "#FFFFFF", bold=True, center=True, max_lines=1)
    return img


SLIDES = {
    4: {
        "title": "ทุกกรณี: เริ่มโปรเจกต์ให้มีเกณฑ์และทะเบียน",
        "subtitle": "Baseline controls ช่วง Initiation ที่ต้องมีตั้งแต่ต้น",
        "tag": "ทุกกรณี | Phase 1",
        "condition": "Applied condition: ทุกกรณี",
        "narrative": "ทุก AI project ต้องบอกให้ชัดว่าอะไรคือผลสำเร็จ ประเมิน risk และลงทะเบียน model/use case ก่อนเดินงานต่อ",
        "steps": [
            ("กำหนด success criteria", "ระบุ AI Model Evaluation Criteria / Business KPI ลงใน BRD"),
            ("ทำ AI Risk Assessment", "กรอก AI Risk Assessment Form และส่ง OpRisk review"),
            ("ลงทะเบียน AI use case/model", "ลงใน Model Inventory ของ AI-COE"),
            ("เก็บ Registration ID", "ใช้เป็น reference ตลอด lifecycle และตอนส่ง evidence"),
        ],
        "evidence": ["BRD", "AI Risk Assessment Form", "Registration ID"],
        "phases": ["Phase 1 Initiation", "OpRisk Review", "AI-COE Inventory"],
    },
    5: {
        "title": "ทุกกรณี: Build / Validate / Data-ARB",
        "subtitle": "สิ่งที่ต้องเตรียมระหว่างออกแบบและพัฒนาโมเดล",
        "tag": "ทุกกรณี | Phase 2",
        "condition": "Applied condition: ทุกกรณี",
        "narrative": "ช่วง build ต้องพิสูจน์ว่าโมเดลอธิบายได้ ทดสอบกับข้อมูลที่ไม่เคยเห็น และตอบคำถาม data governance / architecture review ได้",
        "steps": [
            ("ทำ Explainability section", "ใส่ใน AI Safety & Principles Checklist ให้ผู้เกี่ยวข้องเข้าใจผลลัพธ์"),
            ("ทดสอบ Validation Set", "ใช้ hold-out / unseen data และเทียบกับ criteria ที่กำหนดไว้"),
            ("จัดทำ Model Performance Report", "บันทึกผลทดสอบและข้อจำกัดของโมเดล"),
            ("กรอก ARB Pack ด้าน data", "ระบุ Data Storage Class, Retention, Criticality Tier และ RPO/RTO"),
        ],
        "evidence": ["AI Safety Checklist", "Model Performance Report", "ARB Pack"],
        "phases": ["Phase 2 Development", "Architecture Review"],
    },
    6: {
        "title": "ทุกกรณี: Go-live แล้วต้อง monitor และ secure",
        "subtitle": "Baseline controls ช่วง implementation และ monitoring",
        "tag": "ทุกกรณี | Phase 3-4",
        "condition": "Applied condition: ทุกกรณี",
        "narrative": "ก่อน go-live ต้องพร้อมด้าน resilience และหลัง go-live ต้องมี monitoring, security และ rate limit ที่ตรวจย้อนหลังได้",
        "steps": [
            ("Configure ความพร้อมระบบ", "ตั้ง HA, Backup, DR และ environment tier ตาม Criticality ก่อน go-live"),
            ("ตั้ง Monitoring Pipeline", "ติดตาม drift detection และ performance metrics ตามรอบ review"),
            ("ทำ Monitoring Report", "แชร์ dashboard/report ให้ AI-COE และ OpRisk ตามรอบ"),
            ("ทำ security controls", "Implement encryption, data access control และ rate limit ต่อ user/API key"),
        ],
        "evidence": ["Monitoring Report", "Dashboard", "Architecture Review Record", "ARB Approval / Sign-off", "Security Requirement Checklist"],
        "phases": ["Phase 3 Implementation", "Phase 4 Monitoring"],
    },
    7: {
        "title": "AI Risk = Medium or High: ต้องทดสอบ edge case",
        "subtitle": "ความเสี่ยงระดับกลางขึ้นไปต้องพิสูจน์กรณียาก ไม่ใช่ดูแค่ average performance",
        "tag": "Risk Medium/High",
        "condition": "Applied condition: AI Risk Assessment = Medium or High",
        "narrative": "ถ้า risk ไม่ใช่ low ให้เตรียมชุดทดสอบที่จำลองกรณียากหรือกรณีผิดปกติ แล้วรายงานผลใน Model Performance Report",
        "steps": [
            ("เตรียม edge case dataset", "รวมกรณียาก กรณีผิดปกติ หรือกรณีที่คาดว่าโมเดลอาจพลาด"),
            ("ทดสอบเทียบ criteria", "ใช้ criteria เดียวกับที่กำหนดไว้ใน BRD / evaluation plan"),
            ("บันทึกผลใน report", "ใส่ผล edge case test ใน Model Performance Report"),
        ],
        "evidence": ["Edge case dataset", "Model Performance Report"],
        "phases": ["Phase 2 Development"],
    },
    8: {
        "title": "AI Risk = High: ต้องมี alert, kill switch และแผนหยุดใช้",
        "subtitle": "High-risk AI ต้องถูก monitor แบบมี threshold และมีทางหยุดระบบได้จริง",
        "tag": "Risk High",
        "condition": "Applied condition: AI Risk Assessment = High",
        "narrative": "งานเสี่ยงสูงต้องมีทั้งการทดสอบ robustness ก่อนใช้งาน และ operational control หลัง go-live ถ้าคุณภาพตกต้องรู้ว่าใครรับผิดชอบและหยุดอย่างไร",
        "steps": [
            ("ทดสอบ robustness", "ใช้ noisy / perturbed data และบันทึกผลใน Model Performance Report"),
            ("ตั้ง threshold และ alert", "แจ้งเตือนผ่าน email / Teams / monitoring tool เมื่อ quality ต่ำกว่าเกณฑ์"),
            ("ออกแบบ kill switch", "มี mechanism สำหรับ disable agent หรือหยุดการทำงานของ AI"),
            ("ทำ stop-use plan", "ระบุ owner, response path และแผนหยุดใช้เมื่อเกิด incident"),
        ],
        "evidence": ["Model Performance Report", "Monitoring Dashboard", "Alert configuration", "Incident Response Plan", "Stop-use plan"],
        "phases": ["Phase 2 Development", "Phase 4 Monitoring"],
    },
    9: {
        "title": "งานหลัก / กระทบลูกค้า: ต้องมี human oversight",
        "subtitle": "ใช้กับ AI ที่เกี่ยวข้องกับการตัดสินใจเชิงกลยุทธ์ หรือส่งผลโดยตรงต่อลูกค้า",
        "tag": "Customer Impact",
        "condition": "Applied condition: งานหลักหรือกระทบลูกค้าโดยตรง",
        "narrative": "ถ้า AI มีผลต่อ decision สำคัญหรือลูกค้า ต้องแสดงให้เห็นว่าคนยัง review, approve, override หรือ stop ได้ตาม risk level",
        "steps": [
            ("เลือก oversight model", "กำหนด HITL หรือ HOTL ตาม AI Risk Level"),
            ("ออกแบบ workflow", "ใส่ oversight step ใน Customer Journey / Use Case Diagram / UI flow"),
            ("แสดง control point", "ระบุจุดที่มนุษย์ review, approve, override หรือ stop ระบบได้"),
            ("เก็บภาพเป็นหลักฐาน", "แนบ business process flow, customer journey, use case diagram หรือ UI capture"),
        ],
        "evidence": ["Business process flow", "Customer Journey", "Use Case Diagram", "UI capture"],
        "phases": ["Phase 3 Implementation"],
    },
    10: {
        "title": "Business-critical supervised model: ต้องอธิบายผลลัพธ์ได้",
        "subtitle": "สำหรับโมเดล supervised ที่ใช้กับ credit, risk, compliance หรืองานสำคัญ",
        "tag": "Explainability",
        "condition": "Applied condition: Business-critical supervised model",
        "narrative": "ทีมต้องทำให้ reviewer เห็นว่าโมเดลให้ผลลัพธ์เพราะปัจจัยอะไร ไม่ใช่มีแค่ prediction score",
        "steps": [
            ("เลือกวิธีอธิบายโมเดล", "ใช้ SHAP, LIME หรือ feature importance ตามความเหมาะสม"),
            ("สรุปปัจจัยสำคัญ", "อธิบาย feature ที่มีผลต่อ output ของโมเดล"),
            ("ใส่ใน Model Performance Report", "จัดทำ explainability section ให้ตรวจสอบย้อนหลังได้"),
        ],
        "evidence": ["Model Performance Report", "Explainability section"],
        "phases": ["Phase 2 Development"],
    },
    11: {
        "title": "AI สนทนากับลูกค้า: ต้องเปิดเผยให้ชัด",
        "subtitle": "ใช้กับ chatbot / voicebot / AI agent ที่สื่อสารกับลูกค้าแทนมนุษย์",
        "tag": "Customer-facing AI",
        "condition": "Applied condition: ใช้สนทนาหรือสื่อสารกับลูกค้าแทนมนุษย์",
        "narrative": "ลูกค้าต้องรู้ก่อนใช้งานว่ากำลังคุยกับ AI และข้อมูลถูกใช้เพื่อวัตถุประสงค์อะไร",
        "steps": [
            ("เพิ่ม disclosure บน UI/chatbot", "ระบุชัดว่าระบบกำลังให้บริการด้วย AI"),
            ("แจ้งวัตถุประสงค์การใช้ข้อมูล", "ใช้ wording ที่อ่านง่ายและเห็นก่อนเริ่มบริการ"),
            ("ตรวจจุดแสดงผล", "ให้ disclosure อยู่ในตำแหน่งที่ไม่ถูกซ่อนหรืออ่านยาก"),
            ("เก็บ screenshot", "Capture หน้าจอจริงเป็น evidence ก่อน go-live"),
        ],
        "evidence": ["Screen capture", "UI / chatbot disclosure wording"],
        "phases": ["Phase 3 Implementation"],
    },
    12: {
        "title": "External open data: ต้องรู้ที่มาและ license",
        "subtitle": "Open data ไม่ได้แปลว่าใช้ได้ทุกบริบท ต้องมี provenance ใน ARB Pack",
        "tag": "Open Data",
        "condition": "Applied condition: ใช้งาน External data ซึ่งเป็น Open data",
        "narrative": "ทีมต้องตอบได้ว่าข้อมูลมาจากไหน ใครรวบรวม และ license อนุญาตให้ใช้กับ use case นี้หรือไม่",
        "steps": [
            ("ระบุ data provenance", "บันทึก source, collector, version/date และ license"),
            ("ใส่ใน ARB Pack", "ให้ architecture / governance reviewer ตรวจได้"),
            ("เก็บหลักฐานแหล่งข้อมูล", "แนบ link, snapshot หรือ license evidence ของ source ที่ใช้จริง"),
        ],
        "evidence": ["ARB Pack", "Data provenance record", "License evidence"],
        "phases": ["Phase 2 Development", "Architecture Review"],
    },
}


def render_standard(idx: int, bg: Image.Image, item: dict) -> Image.Image:
    img = soften_background(bg)
    draw = ImageDraw.Draw(img)
    header(draw, item["title"], item["subtitle"], idx, item["tag"])
    badge(draw, item["condition"], 78, 198, C["yellow"], C["ink2"])
    draw_steps_panel(draw, item["steps"])
    draw_right_panels(draw, item["narrative"], item["evidence"], item["phases"])
    return img


def render_handoff(bg: Image.Image) -> Image.Image:
    img = soften_background(bg)
    draw = ImageDraw.Draw(img)
    header(draw, "Evidence handoff: ส่งอะไรให้ gate/reviewer", "ใช้หน้านี้เป็น checklist ปิดท้ายหลังเล่าแต่ละ condition", 13, "HANDOFF")
    rounded(draw, (78, 232, 1040, 890), (255, 255, 255, 240), outline=C["line"], width=2, radius=30)
    draw.rectangle((78, 232, 90, 890), fill=C["yellow"])
    draw_text(draw, "Minimum evidence pack", 122, 268, 860, f(38), C["ink2"], bold=True, max_lines=1)
    packs = [
        ("Initiation", "BRD, AI Risk Assessment Form, Registration ID"),
        ("Development", "AI Safety Checklist, Model Performance Report, Edge case / Robustness result"),
        ("Architecture / Data", "ARB Pack, Architecture Review Record, ARB Approval / Sign-off"),
        ("Go-live / Monitoring", "Monitoring Dashboard/Report, Alert configuration, Security controls"),
        ("Condition-specific", "UI screenshots, Customer Journey, Incident Response / Stop-use plan, Data provenance"),
    ]
    y = 340
    for idx, (head, detail) in enumerate(packs, start=1):
        rounded(draw, (128, y, 988, y + 86), C["soft"], outline="#ECE7E2", radius=18)
        draw.ellipse((150, y + 18, 198, y + 66), fill=C["yellow"])
        draw_text(draw, str(idx), 150, y + 27, 48, f(22), C["ink2"], bold=True, center=True)
        draw_text(draw, head, 220, y + 13, 720, f(25), C["ink2"], bold=True, max_lines=1)
        draw_text(draw, detail, 220, y + 48, 720, f(21), C["muted"], max_lines=1)
        y += 102

    rounded(draw, (1110, 232, 1812, 520), C["taupe"], radius=30)
    draw_text(draw, "วิธีใช้ตอนคุยกับทีม", 1150, 274, 620, f(34), C["yellow"], bold=True, max_lines=1)
    draw_text(draw, "เริ่มจากถามว่า project เข้า condition ไหน จากนั้น assign owner ให้แต่ละ evidence และเช็คก่อนเข้า gate ที่เกี่ยวข้อง", 1152, 342, 610, f(27), "#FFFFFF", max_lines=4, line_gap=7)

    rounded(draw, (1110, 570, 1812, 890), (255, 245, 207, 246), outline=C["line"], width=2, radius=30)
    draw_text(draw, "Control story ที่ควรเล่า", 1150, 612, 620, f(33), C["ink2"], bold=True, max_lines=1)
    story = [
        "1. เงื่อนไขนี้เปิด control อะไร",
        "2. ทีมต้องทำ action ไหนก่อน",
        "3. evidence อยู่ไฟล์ไหน ใครเป็น owner",
        "4. reviewer ตรวจอะไร และต้องส่งเมื่อไร",
    ]
    yy = 678
    for line in story:
        draw_text(draw, line, 1152, yy, 610, f(25), C["ink2"], max_lines=1)
        yy += 48
    return img


def load_bg(idx: int) -> Image.Image:
    path = ORIGINALS / f"slide-{idx:02d}.png"
    if not path.exists():
        raise FileNotFoundError(path)
    return Image.open(path)


def save_slide(idx: int, img: Image.Image):
    FINAL.mkdir(parents=True, exist_ok=True)
    out = FINAL / f"slide-{idx:02d}.png"
    root_copy = IMAGE_ROOT / f"slide-{idx:02d}.png"
    img.convert("RGB").save(out, quality=96)
    img.convert("RGB").save(root_copy, quality=96)


def build_pptx(count: int):
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank = prs.slide_layouts[6]
    for idx in range(1, count + 1):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(FINAL / f"slide-{idx:02d}.png"), 0, 0, width=prs.slide_width, height=prs.slide_height)
    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_PPTX)


def make_contact_sheet(count: int):
    cols = 4
    thumb_w, thumb_h, label_h = 480, 270, 34
    rows = (count + cols - 1) // cols
    sheet = Image.new("RGB", (cols * thumb_w, rows * (thumb_h + label_h)), "#090909")
    d = ImageDraw.Draw(sheet)
    for idx in range(1, count + 1):
        img = Image.open(FINAL / f"slide-{idx:02d}.png").resize((thumb_w, thumb_h), Image.Resampling.LANCZOS).convert("RGB")
        x = ((idx - 1) % cols) * thumb_w
        y = ((idx - 1) // cols) * (thumb_h + label_h)
        d.rectangle((x, y, x + thumb_w, y + label_h), fill="#111111")
        d.text((x + 14, y + 7), f"Slide {idx:02d}", fill="#FFFFFF", font=ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 20))
        sheet.paste(img, (x, y + label_h))
    sheet.save(CONTACT_SHEET)


def write_notes():
    outline = """# AI Control Guideline Workflow

Source files:
- `/Users/nuthdanai/Downloads/AI_Risk_Assessment_v2and_Minimum Control Checklist.xlsx`
- `/Users/nuthdanai/Downloads/AISLC.pptx`

Scope:
- ใช้เฉพาะแถวใน `Control_Guideline` ที่ `Applicable = TRUE`
- รวม 21 applicable controls แปลงเป็น workflow deck 13 slides
- สไลด์เป็น PNG เต็มหน้า มีข้อความฝังในภาพ แล้วนำไปวางเต็มหน้าใน PPTX

Slide outline:
1. Cover: AI Control Guideline Workflow
2. วิธีใช้ deck เป็น workflow
3. Applicable condition map
4. ทุกกรณี: Initiation / Criteria / Risk Assessment / Registration
5. ทุกกรณี: Build / Validate / Data-ARB
6. ทุกกรณี: Go-live / Monitoring / Security
7. AI Risk = Medium or High: Edge case test
8. AI Risk = High: Robustness / Alert / Kill switch / Stop-use plan
9. งานหลักหรือกระทบลูกค้า: Human oversight
10. Business-critical supervised model: Explainability
11. AI สนทนากับลูกค้า: Disclosure + screenshot
12. External open data: Provenance + license
13. Evidence handoff checklist
"""
    OUTLINE.parent.mkdir(parents=True, exist_ok=True)
    OUTLINE.write_text(outline, encoding="utf-8")
    plan = {
        "project": PROJECT,
        "slide_count": 13,
        "mode": "image-generator backgrounds + baked-in Thai text overlay",
        "source_sheet": "Control_Guideline",
        "filter": "Applicable == TRUE",
        "outputs": {
            "pptx": str(OUT_PPTX),
            "images": str(FINAL),
            "contact_sheet": str(CONTACT_SHEET),
            "outline": str(OUTLINE),
        },
    }
    PLAN_JSON.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")


def main():
    renderers = {
        1: lambda bg: render_cover(bg),
        2: lambda bg: render_overview(bg),
        3: lambda bg: render_condition_map(bg),
        13: lambda bg: render_handoff(bg),
    }
    for idx in range(1, 14):
        bg = load_bg(idx)
        if idx in renderers:
            img = renderers[idx](bg)
        else:
            img = render_standard(idx, bg, SLIDES[idx])
        save_slide(idx, img)
    build_pptx(13)
    make_contact_sheet(13)
    write_notes()
    print(OUT_PPTX)
    print(CONTACT_SHEET)
    print(OUTLINE)


if __name__ == "__main__":
    main()
